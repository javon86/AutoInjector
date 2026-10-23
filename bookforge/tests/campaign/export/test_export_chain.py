"""Engine-level export chain (G02 -> G01 -> G05 -> G04 -> G07 -> G11/G12/G13) checked with real validators:
EPUBCheck (java jar), LibreOffice, qpdf, pypdf/pdfium; H09 file checks; determinism; shelf fingerprints;
150k-word performance."""
from __future__ import annotations

import hashlib
import io
import json
import resource
import shutil
import subprocess
import time
import zipfile
from types import SimpleNamespace

import pytest

from bf_export_helpers import (EPUBCHECK_JAR, docx_blocks, epub_blocks, epubcheck_validator, have_epubcheck, in_order,
                               nows, pdf_pages_text, run_epubcheck)
from conftest import PRINCIPAL, Harness, errors_of
from test_set_g import Book, fixture_asset, png_bytes, set_release_profiles
from bookforge import h_checks as hc

CH_A = ("# The Harbour\n" + "\n".join(f"Line {i}: the tide came in over the stones — “slowly,” she said — & the "
                                      f"gulls <wheeled> above the café roofs." for i in range(60))
        + "\n## Night\nThe lamps were lit one by one.\n***\nMorning came grey and cold.")
CH_B = "# The Road\nShe left before dawn.\nThe road ran east along the salt flats."
CH_C = "# Coda\n" + "Z" * 5000 + "\nThe end."

needs_epubcheck = pytest.mark.skipif(not have_epubcheck(), reason=f"EPUBCheck jar not found at {EPUBCHECK_JAR}")
needs_soffice = pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")


class AnyAI:
    """Fake adapter for G08/G09 creative items: returns the canned items for any call."""

    def __init__(self, items):
        self.items = items

    def submit(self, slot, system, user, schema):
        return {"items": self.items}, {"input_tokens": 1, "output_tokens": 1, "model": "fake"}


def full_chain(h, texts=(CH_A, CH_B, CH_C), *, art=False, cover=False):
    """Book with title page/copyright/TOC/colophon (+ optional placed art and EBOOK cover) exported to all
    three formats from the converged G07."""
    b = Book(h, texts=texts)
    set_release_profiles(h, ["MANUSCRIPT_PDF", "EPUB_REFLOWABLE", "EDITABLE_DOCX"], artwork=art)
    g02 = b.g02()
    g01 = b.g01(g02)
    g05 = b.g05(g01)
    inputs = b.matter(g01, g05)
    if art:
        illus = fixture_asset(h, "D07", png_bytes(40, 30, alpha=True))
        placement = {"summary": "art", "links": [illus], "facts": [
            {"subject_id": "pl-1", "predicate": p, "value": v, "assertion_type": "FACT", "effective_from": None,
             "effective_to": None, "source_refs": []}
            for p, v in (("x.g03.anchor_id", f"{b.chapters[0]['record_id']}-p3"), ("x.g03.width_mm", 80),
                         ("x.g03.height_mm", 60), ("x.g03.fit", "CONTAIN"))]}
        g03, _ = b.ok("G03", {"ordered_manuscript_ref": g01, "illustration_policy": "REQUIRED",
                              "placements": [placement]})
        inputs.append(g03)
    g01 = b.g01(g02, input_refs=inputs)
    g05 = b.g05(g01)
    g04 = b.g04(g01, g05)
    g07 = b.g07(g04)
    profiles = {"G11": b.output_profile("MANUSCRIPT_PDF", artwork=art),
                "G12": b.output_profile("EPUB_REFLOWABLE", page=None, artwork=art),
                "G13": b.output_profile("EDITABLE_DOCX", artwork=art)}
    extra = []
    if cover:
        cov = fixture_asset(h, "D11", png_bytes(60, 90))
        g10, _ = b.ok("G10", {"cover_asset_ref": cov, "cover_mode": "EBOOK"})
        extra = [g10]
    b.extra_g12 = extra
    out = {}
    for op in ("G11", "G12", "G13"):
        r = b.export(op, g07, profile=profiles[op], input_refs=extra if op == "G12" else [])
        assert r["status"] == "DRAFT", (op, r["errors"], r["findings"])
        out[op] = (r, b.binary(r["output_refs"][0])[0])
    return b, out


@pytest.fixture(scope="module")
def chain(tmp_path_factory, contracts):
    """One full chain (art + cover + matter) shared by the read-only checks of this module."""
    from bookforge.engine import Engine
    from bookforge.store import Store
    tmp = tmp_path_factory.mktemp("chain")
    store = Store(tmp / "data")
    eng = Engine(store, contracts)
    eng.providers.secret_path = tmp / "secrets.json"
    mh = Harness(eng, eng.bootstrap_project(PRINCIPAL, "Chain Book"))
    yield full_chain(mh, art=True, cover=True)
    store.close()


@pytest.fixture(scope="module")
def chain_epubcheck(chain, tmp_path_factory):
    if not have_epubcheck():
        pytest.skip(f"EPUBCheck jar not found at {EPUBCHECK_JAR}")
    return run_epubcheck(chain[1]["G12"][1], tmp_path_factory.mktemp("epubcheck"))


def chapter_paragraphs(texts):
    from bookforge.exporters.book import parse_chapter
    out = []
    for i, t in enumerate(texts):
        _, blocks = parse_chapter(f"c{i}", t)
        out += [b["text"] for b in blocks if b["t"] in ("p", "h2")]
    return out


# ----------------------------------------------------------------------------- EPUB
def test_epubcheck_only_reports_the_landmark_error(chain_epubcheck):
    """Pins the current EPUBCheck result so any additional error is caught (RSC-011 is NEW-export-1)."""
    rc, msgs = chain_epubcheck
    assert all("RSC-011" in m and "nav.xhtml" in m for m in msgs), msgs


def test_epub_passes_epubcheck(chain_epubcheck):
    rc, msgs = chain_epubcheck
    assert rc == 0 and msgs == [], msgs


@needs_epubcheck
def test_h09_epub_check_with_real_epubcheck(chain, monkeypatch):
    _, out = chain
    monkeypatch.setitem(hc.EXTERNAL_VALIDATORS, "epubcheck", epubcheck_validator)
    issues, info = hc.epub_check(out["G12"][1], run_external=True)
    assert info.validator == "EPUBCheck 5.1.0"
    assert [i for i in issues if i["decision"] == "FAIL"] == []


@needs_epubcheck
def test_epub_from_compiled_manuscript_passes_epubcheck(h, tmp_path):
    b = Book(h, texts=(CH_A, CH_B))
    r = b.export("G12", b.g01(b.g02()))
    assert r["status"] == "DRAFT", r
    rc, msgs = run_epubcheck(b.binary(r["output_refs"][0])[0], tmp_path)
    assert msgs == [], msgs


def test_epub_content_and_structure(chain):
    b, out = chain
    epub = out["G12"][1]
    assert epub_blocks(epub) == COPYRIGHT + chapter_paragraphs((CH_A, CH_B, CH_C)) + COLOPHON
    zf = zipfile.ZipFile(io.BytesIO(epub))
    assert zf.infolist()[0].filename == "mimetype" and zf.infolist()[0].compress_type == zipfile.ZIP_STORED
    from lxml import etree
    opf = etree.fromstring(zf.read("OEBPS/content.opf"))
    ns = {"o": "http://www.idpf.org/2007/opf"}
    ids = [i.get("id") for i in opf.findall(".//o:item", ns)]
    assert len(ids) == len(set(ids))
    hrefs = {i.get("href") for i in opf.findall(".//o:item", ns)}
    packaged = {n[len("OEBPS/"):] for n in zf.namelist() if n.startswith("OEBPS/") and n != "OEBPS/content.opf"}
    assert hrefs == packaged  # every file declared, every declaration packaged
    issues, _ = hc.epub_check(epub, run_external=False)
    assert [i for i in issues if i["decision"] == "FAIL"] == []


MD_ID = "urn:uuid:3f1c2d7e-0000-4000-8000-000000000001"
COPYRIGHT = ["Copyright 2026 Mara Vell. All rights reserved.", MD_ID]
COLOPHON = ["The Salt Road", "Mara Vell", "Copyright 2026 Mara Vell. All rights reserved.", MD_ID]


# ----------------------------------------------------------------------------- PDF
def test_pdf_with_qpdf_pypdf_pdfium(chain, tmp_path):
    b, out = chain
    pdf = out["G11"][1]
    if shutil.which("qpdf"):
        p = tmp_path / "b.pdf"
        p.write_bytes(pdf)
        r = subprocess.run(["qpdf", "--check", str(p)], capture_output=True, text=True)
        assert r.returncode == 0, r.stdout + r.stderr
    from pypdf import PdfReader
    rd = PdfReader(io.BytesIO(pdf), strict=True)
    assert rd.metadata.title == "The Salt Road" and rd.metadata.author == "Mara Vell"
    for pg in rd.pages:
        assert [round(float(v), 2) for v in pg.mediabox] == [0, 0, 432.0, 648.0]  # 6 x 9 in
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(pdf)
    for i in range(len(doc)):  # every page renders and carries ink
        im = doc[i].render(scale=0.5).to_pil().convert("L")
        assert im.getextrema()[0] < 128, f"page {i + 1} renders blank"
    text = "".join(pdf_pages_text(pdf, bottom=53))  # G04 margins are 54 pt; footer below
    assert in_order(text, chapter_paragraphs((CH_A, CH_B, CH_C))) == []
    assert hc.pdf_structure(pdf)[1] == [] or all(i["decision"] == "PASS" for i in hc.pdf_structure(pdf)[1])


def test_pdf_toc_pages_match_g07(chain):
    b, out = chain
    h = b.h
    r11 = out["G11"][0]
    import pypdfium2 as pdfium
    pdf = out["G11"][1]
    doc = pdfium.PdfDocument(pdf)
    # find the converged G07 entries through the export's provenance (the G07 descriptor is pinned)
    bin11 = b.binary(r11["output_refs"][0])[2]
    g07 = next(h.record(r) for r in bin11["source_refs"] if r["module_id"] == "G07")
    entries = json.loads(b.binary(g07["ref"])[0])["entries"]
    for e in entries:
        page_text = nows(doc[e["page"] - 1].get_textpage().get_text_range())
        assert nows(e["title"])[:40] in page_text, (e["title"], e["page"])


# ----------------------------------------------------------------------------- DOCX
def test_docx_python_docx_reads_structure(chain):
    b, out = chain
    docx = out["G13"][1]
    assert docx_blocks(docx) == COPYRIGHT + chapter_paragraphs((CH_A, CH_B, CH_C)) + COLOPHON
    import docx as pydocx
    d = pydocx.Document(io.BytesIO(docx))
    sec = d.sections[0]
    assert abs(sec.page_width.mm - 152.4) < 0.1 and abs(sec.page_height.mm - 228.6) < 0.1
    assert len(d.inline_shapes) == 1
    issues, info = hc.docx_check(docx)
    assert [i for i in issues if i["decision"] == "FAIL"] == []
    assert set(info.anchors) <= info.bookmarks


@needs_soffice
def test_docx_opens_in_libreoffice_with_all_text(chain, tmp_path):
    b, out = chain
    (tmp_path / "b.docx").write_bytes(out["G13"][1])
    r = subprocess.run(["soffice", f"-env:UserInstallation=file://{tmp_path}/lo", "--headless", "--convert-to", "pdf",
                        "--outdir", str(tmp_path), str(tmp_path / "b.docx")], capture_output=True, text=True,
                       timeout=180)
    pdf = (tmp_path / "b.pdf").read_bytes()
    text = "".join(pdf_pages_text(pdf))
    assert in_order(text, chapter_paragraphs((CH_A, CH_B, CH_C))) == [], r.stdout + r.stderr
    from pypdf import PdfReader
    rd = PdfReader(io.BytesIO(pdf))
    assert abs(float(rd.pages[0].mediabox.width) - 432) < 1


# ----------------------------------------------------------------------------- determinism / shelf
def test_exports_with_art_and_matter_are_byte_identical_on_rerun(chain):
    b, first = chain
    # re-export from the same converged G07 (pinned by the first export's binary)
    bin11 = b.binary(first["G11"][0]["output_refs"][0])[2]
    g07_ref = next(r for r in bin11["source_refs"] if r["module_id"] == "G07")
    for op, prof in (("G11", b.output_profile("MANUSCRIPT_PDF", artwork=True)),
                     ("G12", b.output_profile("EPUB_REFLOWABLE", page=None, artwork=True)),
                     ("G13", b.output_profile("EDITABLE_DOCX", artwork=True))):
        r = b.export(op, g07_ref, profile=prof, input_refs=b.extra_g12 if op == "G12" else [])
        assert r["status"] == "DRAFT", r
        again = b.binary(r["output_refs"][0])[0]
        assert again == first[op][1], op
        assert r["result"]["sha256"] == hashlib.sha256(again).hexdigest()


def test_shelf_fingerprints_match_files_and_detect_tampering(h):
    from bookforge.shelf import finished_books
    b, out = full_chain(h)
    exports = {"MANUSCRIPT_PDF": {"descriptor": out["G11"][0]["output_refs"][0]},
               "EPUB_REFLOWABLE": {"descriptor": out["G12"][0]["output_refs"][0]},
               "EDITABLE_DOCX": {"descriptor": out["G13"][0]["output_refs"][0]}}
    h.store.set_setting(f"book_path:{h.project_id}:{h.branch_id}", {"exports": exports})
    app = SimpleNamespace(store=h.store)
    books = [x for x in finished_books(app) if x["project_id"] == h.project_id]
    assert len(books) == 1 and books[0]["all_verified"] is True
    by_fmt = {f["format"]: f for f in books[0]["files"]}
    for fmt, op in (("PDF", "G11"), ("EPUB", "G12"), ("DOCX", "G13")):
        data = out[op][1]
        assert by_fmt[fmt]["sha256"] == hashlib.sha256(data).hexdigest()
        assert by_fmt[fmt]["bytes"] == len(data)
    # tamper with the stored EPUB bytes: the shelf must show it as not verified
    rec = h.store.get_record(h.project_id, h.branch_id, by_fmt["EPUB"]["record_id"], by_fmt["EPUB"]["version"])
    path = h.store.asset_dir / rec["binary"]["storage_uri"].split("/")[-1]
    path.write_bytes(out["G12"][1][:-10] + b"tampered!!")
    books = [x for x in finished_books(app) if x["project_id"] == h.project_id]
    by_fmt = {f["format"]: f for f in books[0]["files"]}
    assert by_fmt["EPUB"]["verified"] is False and by_fmt["PDF"]["verified"] is True
    assert books[0]["all_verified"] is False


# ----------------------------------------------------------------------------- matter item IDs
def test_non_ascii_matter_item_id_exports_pdf(h):
    b = Book(h)
    h.engine.providers.test_adapter = AnyAI([{"item_id": "dédicace", "title": "Dédicace", "text": "Pour Marie."}])
    h.engine.providers.configure_slot("AI-1", {"adapter_id": "anthropic"})
    g02 = b.g02()
    g01 = b.g01(g02)
    r = h.create("G08", {"book_ref": g01, "front_matter_items": ["title_page", "dédicace", "toc"],
                         "publication_profile_ref": b.pdf_profile})
    assert r["status"] == "DRAFT", r
    g01 = b.g01(g02, input_refs=[r["output_refs"][0]])
    g05 = b.g05(g01)
    g07 = b.g07(b.g04(g01, g05))
    assert b.export("G12", g07)["status"] == "DRAFT"
    r11 = b.export("G11", g07)
    assert r11["status"] == "DRAFT", errors_of(r11)


# ----------------------------------------------------------------------------- performance
def test_150k_word_book_exports_quickly(h):
    import random
    rng = random.Random(1)
    words = ("the tide came in over stones slowly she said gulls wheeled above café roofs while ferry waited "
             "“quote” it’s").split()
    texts = [f"# Chapter {c}\n" + "\n".join(" ".join(rng.choice(words) for _ in range(62)) for _ in range(60))
             for c in range(40)]
    assert 140_000 < sum(len(t.split()) for t in texts) < 160_000
    b = Book(h, texts=texts)
    t0 = time.perf_counter()
    g02 = b.g02()
    g01 = b.g01(g02)
    g05 = b.g05(g01)
    g07 = b.g07(b.g04(g01, g05))
    sizes = {}
    for op in ("G11", "G12", "G13"):
        r = b.export(op, g07)
        assert r["status"] == "DRAFT", r
        sizes[op] = len(b.binary(r["output_refs"][0])[0])
    elapsed = time.perf_counter() - t0
    assert elapsed < 30, f"assembly + 3 exports took {elapsed:.1f}s"
    assert resource.getrusage(resource.RUSAGE_SELF).ru_maxrss < 1_500_000  # KiB: < ~1.5 GB peak for the process
    assert sizes["G11"] < 10_000_000 and sizes["G12"] < 5_000_000 and sizes["G13"] < 5_000_000
