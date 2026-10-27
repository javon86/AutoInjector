"""🎨 B3 book covers: the design API (defaults, profile fall-backs, validation, revisions, blurb / bio in the profile),
the templates (short and very long titles: every glyph inside the safe area, read back from the PDF itself), the
print maths (spine width for several page counts and papers, wrap size with bleed, page-count checks, spine text only
from 100 pages, rotated and centred, barcode box), the AI art job through a fake AutoInjector bridge (purpose
"cover", size, upscale, character lock and the NO_LOCK fallback, every error code), upload, the ebook cover (PDF /
SVG / art), the book's PDF / EPUB / DOCX with the cover, the book save and the HTTP API."""
from __future__ import annotations

import io
import json
import re
import threading
import zipfile
from http.server import ThreadingHTTPServer

import docx as python_docx
import pypdf
import pypdfium2 as pdfium
import pytest
from PIL import Image

from akowe import cover, forge
from akowe.book_path import get_profile, save_profile
from akowe.errors import AkoweError
from akowe.exporters import book as bookmod
from akowe.exporters import cover as cx
from akowe.exporters.docx import build_docx, inspect_docx
from akowe.exporters.epub import build_epub, inspect_epub
from akowe.exporters.images import load_image
from akowe.exporters.layout import Typography, make_geometry, paginate
from akowe.exporters.pdf import inspect_pdf, render_pdf

from test_failover_review import SlotFake, make_app, no_provider_sleep  # noqa: F401
from test_forge import FakeImageBridge, make_char, png
from test_server import call

BR = "main"
LONG_TITLE = ("The Extraordinary and Entirely Unreliable Chronicle of the Glass Harbour Lighthouse Keepers and "
              "Their Seven Stolen Keys")
TEXTS = {"title": "Salt", "author": "Ada Vey", "subtitle": "A novel", "series": "The Harbour Cycle, Book One",
         "tagline": "Every light hides a debt.",
         "blurb": "When the lights go out, Mara has one night to find the ledger. " * 3,
         "author_bio": "Ada Vey lives by the sea."}


@pytest.fixture
def fx(tmp_path):
    app = make_app(tmp_path, SlotFake(), slots=("AI-1",))
    pid = app.engine.bootstrap_project("author", "Cover book")
    save_profile(app.store, pid, {"title": "The Salt Road", "author": "Mara Vell", "genre": "Fantasy",
                                  "idea": "A courier crosses a salt desert to deliver a key."})
    return app, pid


def wait(app, job):
    return forge.wait(app, "author", job["job_id"])


# ====================================================================== reading text back from a PDF
def pdf_texts(data: bytes) -> list[dict]:
    """Every text show of every page, read back with pypdf: {text, x, y, size, font, rot, page}."""
    out = []
    reader = pypdf.PdfReader(io.BytesIO(data))
    for pno, page in enumerate(reader.pages):
        fonts = {}
        for name, ref in (page["/Resources"].get("/Font") or {}).items():
            fonts[name] = str(ref.get_object()["/BaseFont"])[1:]
        content = page.get_contents().get_data().decode("latin-1")
        for m in re.finditer(r"BT /(\S+) ([\d.]+) Tf (-?[\d.]+) Tc [\d. ]+ rg (-?[\d.]+) (-?[\d.]+) (-?[\d.]+) "
                             r"(-?[\d.]+) (-?[\d.]+) (-?[\d.]+) Tm \((.*?)\) Tj ET", content):
            a, b, _c, _d, e, f = (float(m.group(i)) for i in range(4, 10))
            raw = re.sub(r"\\(\d{3})", lambda mm: chr(int(mm.group(1), 8)), m.group(10))
            raw = re.sub(r"\\(.)", r"\1", raw)
            text = raw.encode("latin-1").decode("cp1252")
            out.append({"text": text, "font": fonts["/" + m.group(1)], "size": float(m.group(2)),
                        "track": float(m.group(3)), "x": e, "y": f, "rot": 0 if a == 1 else -90, "page": pno})
    return out


def extent(t: dict) -> tuple:
    """(x0, y0, x1, y1) of a text show: the advance width with tracking, ascent and descent of the font."""
    w = cx.measure(t["text"], t["font"], t["size"], t["track"])
    up, down = cx.asc(t["font"]) * t["size"], cx.desc(t["font"]) * t["size"]
    if t["rot"] == 0:
        return t["x"], t["y"] - down, t["x"] + w, t["y"] + up
    return t["x"] - down, t["y"] - w, t["x"] + up, t["y"]


def inside(box, area, tol=0.01) -> bool:
    x0, y0, x1, y1 = box
    ax, ay, aw, ah = area
    return x0 >= ax - tol and y0 >= ay - tol and x1 <= ax + aw + tol and y1 <= ay + ah + tol


def render_front(design, art=None, size=cx.EBOOK_PT):
    W, H = size
    ops, info = cx.layout_front(design, (0, 0, W, H), art=art)
    pdf = cx.render_pdf([{"width": W, "height": H, "ops": ops}], images={"art": art} if art else {})
    return pdf, ops, info


# ====================================================================== templates
@pytest.mark.parametrize("tid", list(cx.TEMPLATES))
@pytest.mark.parametrize("title", ["Salt", LONG_TITLE])
def test_every_template_keeps_text_in_the_safe_area(tid, title):
    art = load_image(png(400, 600, (40, 90, 140)))
    for with_art in (art, None):
        d = {"template": tid, "texts": dict(TEXTS, title=title)}
        pdf, ops, info = render_front(d, with_art)
        assert not inspect_pdf(pdf)["errors"]
        shows = pdf_texts(pdf)
        assert shows, tid
        safe = info["safe"]
        for t in shows:
            assert inside(extent(t), safe), (tid, t["text"], extent(t), safe)
        # the read-back positions are the layout's own
        laid = [op for op in ops if op[0] == "text"]
        assert len(laid) == len(shows)
        for op, t in zip(laid, shows):
            assert abs(op[4] - t["x"]) < 0.01 and abs(op[5] - t["y"]) < 0.01
        title_lines = [t for op, t in zip(laid, shows) if op[10] == "title"]
        assert all(t["size"] == max(s["size"] for s in shows) for t in title_lines)  # the title leads
        joined = " ".join(t["text"] for t in title_lines).replace("…", "")
        assert joined.split()[0].lower() == title.split()[0].lower()
        if title == "Salt":
            assert len(info["title_lines"]) == 1 and info["title_size"] >= 0.11 * cx.EBOOK_PT[0]
        else:
            assert 3 <= len(info["title_lines"]) <= 5 and info["title_size"] >= 14
            widths = [cx.measure(t["text"], t["font"], t["size"], t["track"]) for t in title_lines]
            assert min(widths) > 0.33 * max(widths)  # balanced line breaks: no lonely last word
        for n in ("ADA VEY", "Ada Vey"):
            if any(t["text"] == n for t in shows):
                break
        else:
            pytest.fail(f"{tid}: the author is missing")


def test_title_fitting_shrinks_and_breaks_long_titles():
    w = 300.0
    s, lines, cut = cx.fit_text("Salt", "Times-Bold", max_size=60, min_size=14, width=w, max_lines=3, max_height=200)
    assert s == 60 and lines == ["Salt"] and not cut
    s, lines, cut = cx.fit_text(LONG_TITLE, "Times-Bold", max_size=60, min_size=14, width=w, max_lines=4,
                                max_height=200)
    assert s < 60 and 2 <= len(lines) <= 4 and not cut
    assert all(cx.measure(x, "Times-Bold", s) <= w + 1e-6 for x in lines)
    lens = [cx.measure(x, "Times-Bold", s) for x in lines]
    assert min(lens) > 0.45 * max(lens)  # balanced, no lonely last word
    # an absurd title at the minimum size is cut with an ellipsis
    s, lines, cut = cx.fit_text("word " * 200, "Times-Bold", max_size=30, min_size=14, width=w, max_lines=3,
                                max_height=80)
    assert cut and lines[-1].endswith("…") and s == 14
    # a word longer than the line is broken, never overflows
    s, lines, _ = cx.fit_text("Supercalifragilisticexpialidocious" * 3, "Helvetica-Bold", max_size=40, min_size=14,
                              width=120, max_lines=4, max_height=300)
    assert all(cx.measure(x, "Helvetica-Bold", s) <= 120 + 1e-6 for x in lines)


def test_contrast_scrim_and_colours():
    art = load_image(png(400, 600))
    d = {"template": "gilded", "texts": TEXTS}
    _pdf, ops, _ = render_front(d, art)
    grads = [op for op in ops if op[0] == "grad"]
    assert grads and all(max(a for _t, a in g[6]) >= 0.8 for g in grads)  # solid enough where the text is
    title = next(op for op in ops if op[0] == "text" and op[10] == "title")
    top_grad = max(grads, key=lambda g: g[2] + g[4])
    assert top_grad[2] <= title[5] <= top_grad[2] + top_grad[4]
    # no artwork: no scrim; colour overrides are used; a bad colour is ignored by the layout
    _pdf, ops, _ = render_front(dict(d, colors={"background": "#000001", "text": "#ff0000", "accent": "nope"}))
    assert not [op for op in ops if op[0] == "grad"]
    assert ops[0][5] == (0, 0, 1) and next(op for op in ops if op[0] == "text" and op[10] == "title")[6] == (255, 0, 0)
    pdf, _ops, _ = render_front(dict(d, colors={"background": "#000001"}))
    assert b"0 0 0.004 rg" in pdf  # near-black stays near-black (never read as 0-1 floats)


def test_rendered_front_cover_has_the_art_and_dark_scrim():
    art = load_image(png(400, 600, (250, 250, 250)))
    pdf, _ops, _ = render_front({"template": "thriller", "texts": TEXTS}, art)
    page = pdfium.PdfDocument(pdf)[0]
    img = page.render(scale=0.5).to_pil().convert("RGB")
    w, h = img.size
    mid = img.getpixel((w // 2, h // 3))
    low = img.getpixel((4, h - 4))
    assert sum(mid) > 600  # the white art shows in the middle
    assert sum(low) < 200  # the scrim darkens the art behind the author


def test_svg_matches_the_layout_and_pins_widths():
    art = load_image(png(400, 600))
    W, H = cx.EBOOK_PT
    ops, _ = cx.layout_front({"template": "fantasy", "texts": TEXTS}, (0, 0, W, H), art=art)
    svg = cx.render_svg(W, H, ops, hrefs={"art": "data:image/png;base64,AAAA"}, title="Cover: Salt")
    import xml.etree.ElementTree as ET
    root = ET.fromstring(svg)
    ns = "{http://www.w3.org/2000/svg}"
    texts = root.findall(f".//{ns}text")
    assert len(texts) == len([o for o in ops if o[0] == "text"])
    for t in texts:
        assert float(t.get("textLength")) > 0 and t.get("lengthAdjust") == "spacingAndGlyphs"
    assert root.findall(f".//{ns}image") and root.findall(f".//{ns}linearGradient")


# ====================================================================== print maths
@pytest.mark.parametrize("pages,paper,spine", [(24, "white", 0.054048), (100, "white", 0.2252), (101, "cream", 0.255),
                                               (320, "cream", 0.8), (500, "white", 1.126), (776, "cream", 1.94)])
def test_spine_width_from_page_count_and_paper(pages, paper, spine):
    n, notes = cx.check_print("6x9", paper, pages)
    g = cx.wrap_geometry("6x9", paper, n)
    assert abs(g["spine_in"] - spine) < 1e-4, g
    assert bool(notes) == (pages % 2 == 1)
    assert abs(g["width_in"] - (0.125 + 6 + g["spine_in"] + 6 + 0.125)) < 1e-4 and g["height_in"] == 9.25


def test_page_count_checks_are_plain():
    for trim, paper, pages, words in (("7x10", "white", 200, "trim size"), ("6x9", "glossy", 200, "white or cream"),
                                      ("6x9", "white", None, "page count"), ("6x9", "white", 12, "at least 24"),
                                      ("6x9", "cream", 800, "at most 776"), ("6x9", "white", "many", "whole number")):
        with pytest.raises(cx.CoverError) as e:
            cx.check_print(trim, paper, pages)
        assert words in str(e.value)


@pytest.mark.parametrize("trim", list(cx.TRIMS))
def test_wrap_pdf_size_bleed_spine_and_barcode(trim):
    art = load_image(png(300, 450))
    d = {"template": "gilded", "texts": TEXTS}
    g = cx.wrap_geometry(trim, "cream", 300)
    ops, info = cx.layout_wrap(d, g, art=art)
    p = g["pt"]
    b = p["bleed"]
    pdf = cx.render_pdf([{"width": p["width"], "height": p["height"], "ops": ops,
                          "trim": (b, b, p["width"] - 2 * b, p["height"] - 2 * b)}], images={"art": art})
    rep = inspect_pdf(pdf)
    assert not rep["errors"]
    tw, th = cx.TRIMS[trim]
    mb = rep["media_boxes"][0]
    assert abs(mb[2] - (2 * tw + 300 * 0.0025 + 0.25) * 72) < 0.01 and abs(mb[3] - (th + 0.25) * 72) < 0.01
    reader = pypdf.PdfReader(io.BytesIO(pdf))
    tb = [float(v) for v in reader.pages[0].trimbox]
    assert tb == pytest.approx([9.0, 9.0, mb[2] - 9.0, mb[3] - 9.0], abs=0.01)
    shows = pdf_texts(pdf)
    spine = [t for t in shows if t["rot"] == -90]
    assert [t["text"] for t in spine] == ["THE SALT ROAD".replace("THE SALT ROAD", "SALT"), "ADA VEY"]
    sx0, sw = p["spine_x"], p["spine"]
    m = cx.SPINE_MARGIN_IN * 72
    for t in spine:
        x0, y0, x1, y1 = extent(t)
        assert x0 >= sx0 + m - 0.01 and x1 <= sx0 + sw - m + 0.01
        assert y0 >= b + 36 - 0.01 and y1 <= b + p["trim_h"] - 36 + 0.01
    # centred across the spine on the capital height, and along it
    tt = spine[0]
    assert abs((tt["x"] + cx.cap(tt["font"]) * tt["size"] / 2) - (sx0 + sw / 2)) < 0.01
    first_top, last_bottom = spine[0]["y"], extent(spine[-1])[1]
    assert abs((first_top + last_bottom) / 2 - (b + p["trim_h"] / 2)) < 0.5
    # back and front texts stay in their panels' safe areas; the barcode box is clear
    bx, by, bw, bh = info["back"]["barcode"]
    assert abs(bw - 144) < 0.01 and abs(bh - 86.4) < 0.01
    assert abs((bx + bw) - (sx0 - 18)) < 0.01 and abs(by - (b + 18)) < 0.01
    for t in shows:
        if t["rot"]:
            continue
        box = extent(t)
        area = info["back"]["safe"] if box[0] < sx0 else info["front"]["safe"]
        assert inside(box, area), (t["text"], box, area)
        assert not (box[0] < bx + bw and box[2] > bx and box[1] < by + bh and box[3] > by), t["text"]
    assert any(t["text"].startswith("When the lights go out") for t in shows)
    assert any(t["text"] == "ABOUT THE AUTHOR" for t in shows)
    # the barcode box is white in the rendered file
    img = pdfium.PdfDocument(pdf)[0].render(scale=1).to_pil().convert("RGB")
    assert img.getpixel((int(bx + bw / 2), int(img.size[1] - (by + bh / 2)))) == (255, 255, 255)


def test_no_spine_text_under_100_pages_and_long_blurbs_are_cut():
    d = {"template": "thriller", "texts": dict(TEXTS, blurb="A very long blurb sentence goes here. " * 200)}
    g = cx.wrap_geometry("5x8", "white", 98)
    ops, info = cx.layout_wrap(d, g)
    assert not [op for op in ops if op[0] == "text" and op[8]]
    assert any("spine text needs at least 100 pages" in w for w in info["warnings"])
    assert any("too long for the back cover" in w for w in info["warnings"])
    floor = info["back"]["floor"]
    back = [op for op in ops if op[0] == "text" and op[4] < g["pt"]["spine_x"]]
    assert back and min(op[5] - cx.desc(op[2]) * op[3] for op in back) >= floor - 0.01
    assert any(op[1].endswith("…") for op in back)
    g = cx.wrap_geometry("5x8", "white", 100)
    ops, info = cx.layout_wrap(d, g)
    assert [op for op in ops if op[0] == "text" and op[8]] and info["spine"]["spine_text"]


# ====================================================================== the design API
def test_design_defaults_profile_fallbacks_validation_and_revisions(fx):
    app, pid = fx
    out = cover.public(app, pid, BR)
    d = out["design"]
    assert d["template"] == "gilded" and d["revision"] == 0 and d["art"] is None
    assert out["texts"]["title"] == "The Salt Road" and out["texts"]["author"] == "Mara Vell"
    assert len(out["templates"]) >= 6 and {t["light"] for t in out["templates"]} == {True, False}
    assert out["print"]["ok"] is False and "page count" in out["print"]["problem"]
    assert out["rules"]["spine_text_min_pages"] == 100
    out = cover.save(app, pid, BR, {"template": "scifi", "texts": {"title": "Salt", "subtitle": "A novel"},
                                    "colors": {"accent": "#FFAA00"}, "blurb": "  The back text. ",
                                    "author_bio": "Born at sea.", "print": {"trim": "5x8", "paper": "white",
                                                                            "page_count": 240}})
    d = out["design"]
    assert d["revision"] == 1 and d["template"] == "scifi" and d["colors"] == {"accent": "#ffaa00"}
    assert out["texts"]["title"] == "Salt" and out["texts"]["author"] == "Mara Vell"  # the profile's author
    assert out["texts"]["blurb"] == "The back text." and get_profile(app.store, pid)["author_bio"] == "Born at sea."
    assert out["print"]["ok"] and out["print"]["spine_in"] == pytest.approx(0.54048, abs=1e-4) and \
        out["print"]["page_count_source"] == "yours"
    with pytest.raises(AkoweError) as e:
        cover.save(app, pid, BR, {"template": "gilded", "revision": 0})
    assert e.value.code == "VERSION_CONFLICT"
    assert cover.save(app, pid, BR, {"colors": {"accent": None}, "revision": 1})["design"]["colors"] == {}
    for bad in ({"template": "comic"}, {"colors": {"text": "red"}}, {"colors": {"glow": "#ffffff"}},
                {"texts": {"title": 5}}, {"texts": {"blurb": "x"}}, {"texts": {"title": "x" * 201}},
                {"print": {"trim": "7x10"}}, {"print": {"paper": "glossy"}}, {"print": {"page_count": 3}},
                {"print": {"page_count": "300"}}, {"use_art": "yes"}, {"art": {"asset_id": "x"}},
                {"blurb": "x" * 3001}, {"nope": 1}, "text"):
        with pytest.raises(AkoweError) as e:
            cover.save(app, pid, BR, bad)
        assert e.value.code == "INVALID_INPUT", bad


def test_page_count_comes_from_the_pdf_export_then_an_estimate(fx, monkeypatch):
    app, pid = fx
    built = book_files()
    monkeypatch.setattr(cover, "_export_bytes", lambda a, p, b, fmt: (built[fmt], {"title": "The Salt Road"})
                        if fmt in built else None)
    n = inspect_pdf(built["PDF"])["page_count"]
    info = cover.public(app, pid, BR)["print"]
    assert info["page_count"] == n and info["page_count_source"] == "pdf"
    # the PDF has fewer than 24 pages: a plain problem, and the request's own number is used
    assert info["ok"] is False and "at least 24" in info["problem"]
    data, mime, name, geo = cover.print_cover(app, pid, BR, pages="120", paper="white", trim="5.5x8.5")
    assert mime == "application/pdf" and geo["page_count"] == 120 and geo["spine_in"] == pytest.approx(0.27024, abs=1e-4)
    assert name.endswith("print cover 5.5x8.5 white 120p.pdf")
    monkeypatch.setattr(cover, "_export_bytes", lambda *a: None)
    monkeypatch.setattr(cover, "_export_files", lambda *a: ([], {"words": 84000}))
    info = cover.public(app, pid, BR)["print"]
    assert info["page_count_source"] == "estimate" and info["page_count"] == 308 and info["ok"]
    assert any("estimate" in n_ for n_ in info["notes"])
    with pytest.raises(AkoweError) as e:
        cover.print_cover(app, pid, BR, pages="12")
    assert "at least 24" in str(e.value)


# ====================================================================== AI art (fake AutoInjector bridge)
@pytest.fixture
def ib(fx):
    app, pid = fx
    b = FakeImageBridge()
    forge.save_bridge_settings(app, {"url": b.url, "token": "img-token-1"})
    yield app, pid, b
    b.close()


def test_art_job_asks_for_a_cover_and_stores_it(ib):
    app, pid, b = ib
    cover.save(app, pid, BR, {"art_note": "a lone rider at dusk", "character_lock": "Mara Vell",
                              "print": {"trim": "5x8"}})
    j = wait(app, app.cover.start_art("author", pid, BR, {"engine": "sd", "upscale": 2}))
    assert j["status"] == "done", j
    r = b.requests[0]
    assert r["purpose"] == "cover" and r["engine"] == "sd" and r["upscale"] == 2 and r["returnImage"] is True
    assert r["size"] == {"width": 1600, "height": 2560} and r["characterLock"] == "Mara Vell"
    assert r["from"] == "akowe-cover" and r["book"] == "The Salt Road"
    assert "a lone rider at dusk" in r["prompt"] and "No text" in r["prompt"] and "Fantasy" in r["prompt"]
    assert b.auth[0] == "Bearer img-token-1"
    out = cover.public(app, pid, BR)
    art = out["design"]["art"]
    assert art["engine"] == "sd" and art["upscale"] == 2 and art["width"] == 60 and art["url"].endswith("/1")
    assert any("300 dpi" in w for w in out["warnings"])  # a 60 px picture is far too small for print
    data, mime, _ = cover.ebook(app, pid, BR, fmt="art")
    assert data == png() and mime == "image/png"
    # ChatGPT, no upscale, no lock: NO_LOCK is retried without the lock
    cover.save(app, pid, BR, {"character_lock": "Nobody"})
    b.script[2] = (400, {"ok": False, "error": "NO_LOCK", "message": "no lock"})
    j = wait(app, app.cover.start_art("author", pid, BR, {"engine": "chatgpt"}))
    assert j["status"] == "done" and "characterLock" in b.requests[1] and "characterLock" not in b.requests[2]
    assert "upscale" not in b.requests[2]
    assert "no character lock" in " ".join(cover.public(app, pid, BR)["design"]["art"]["warnings"])
    for bad in ({"engine": "dalle"}, {"upscale": 3}, {"upscale": True}):
        with pytest.raises(AkoweError):
            app.cover.start_art("author", pid, BR, bad)


@pytest.mark.parametrize("code", ["SD_NOT_CONNECTED", "CHATGPT_NOT_READY", "IMAGE_REFUSED", "IMAGE_TIMEOUT",
                                  "IMAGE_LIMIT", "NO_IMAGE_ENGINE"])
def test_art_error_codes_are_plain(ib, code):
    app, pid, b = ib
    b.script[1] = (503, {"ok": False, "error": code, "message": "short reason", "hint": "do this"})
    j = wait(app, app.cover.start_art("author", pid, BR, {"engine": "auto"}))
    assert j["status"] == "error" and j["error"]["code"] == code and j["error"]["fix"] == "do this"
    assert cover.public(app, pid, BR)["design"]["art"] is None


def test_art_bad_answer_wrong_key_and_unreachable(ib):
    app, pid, b = ib
    b.script[1] = (200, {"ok": True, "engine": "sd", "images": []})
    assert wait(app, app.cover.start_art("author", pid, BR, {}))["error"]["code"] == "IMAGE_BAD_ANSWER"
    b.script[2] = (401, {"ok": False, "error": "UNAUTHORIZED"})
    j = wait(app, app.cover.start_art("author", pid, BR, {}))
    assert j["error"]["code"] == "IMAGE_BRIDGE_KEY" and "Settings" in j["error"]["fix"]
    b.close()
    j = wait(app, app.cover.start_art("author", pid, BR, {}))
    assert j["error"]["code"] == "IMAGE_BRIDGE_UNREACHABLE"
    b.close = lambda: None


def test_upload_remove_and_typographic_cover(fx):
    app, pid = fx
    out = app.cover.upload("author", pid, BR, png(800, 1200, fmt="JPEG"))
    assert out["design"]["art"]["source"] == "upload" and out["design"]["art"]["mime"] == "image/jpeg"
    pdf, mime, name = cover.ebook(app, pid, BR)
    assert mime == "application/pdf" and name == "The Salt Road - ebook cover.pdf"
    assert b"/Subtype /Image" in pdf and not inspect_pdf(pdf)["errors"]
    svg, mime, _ = cover.ebook(app, pid, BR, fmt="svg")
    assert mime == "image/svg+xml" and b"data:image/jpeg;base64," in svg and b">THE SALT<" in svg
    for bad in (b"GIF89a....", b"", "text"):
        with pytest.raises(AkoweError):
            app.cover.upload("author", pid, BR, bad)
    out = cover.save(app, pid, BR, {"use_art": False})
    pdf, _m, _n = cover.ebook(app, pid, BR)
    assert b"/Subtype /Image" not in pdf and out["design"]["art"]  # kept, just not used
    out = cover.save(app, pid, BR, {"use_art": True, "art": None})
    assert out["design"]["art"] is None
    with pytest.raises(AkoweError):
        cover.ebook(app, pid, BR, fmt="art")


# ====================================================================== the book's own files with the cover
def book_files() -> dict:
    meta = {"title": "The Salt Road", "author": "Mara Vell", "identifier": "urn:uuid:1", "language": "en",
            "rights": "Copyright Mara Vell"}
    book = bookmod.build_book(meta, [{"id": "c1", "text": "# One\nThe salt was white. " * 30},
                                     {"id": "c2", "text": "# Two\nThe road was long."}], toc=True)
    page = {"width_mm": 152.4, "height_mm": 228.6, "bleed_mm": 0, "safe_margin_mm": 12.7}
    geom = make_geometry(page, {k: 54 for k in ("top", "bottom", "left", "right")})
    lay = paginate(book, geom, Typography("times", 11.0, 14.0))
    nav = [dict(e, page=None) for e in bookmod.heading_entries(book, 2)]
    return {"PDF": render_pdf(lay, page_width=geom.width, page_height=geom.height, meta=meta),
            "EPUB": build_epub(book, modified="2026-01-01T00:00:00Z", toc_entries=nav),
            "DOCX": build_docx(book, created="2026-01-01T00:00:00Z", toc_entries=nav, page=page)}


@pytest.fixture
def exported(fx, monkeypatch):
    app, pid = fx
    built = book_files()
    monkeypatch.setattr(cover, "_export_bytes", lambda a, p, b, fmt: (built[fmt], {"title": "The Salt Road"}))
    return app, pid, built


@pytest.mark.parametrize("with_art", [True, False])
def test_pdf_export_gets_a_cover_page(exported, with_art):
    app, pid, built = exported
    if with_art:
        app.cover.upload("author", pid, BR, png(300, 450))
    data, mime, name = cover.book_with_cover(app, pid, BR, "PDF")
    assert mime == "application/pdf" and name == "The Salt Road (with cover).pdf"
    before, after = inspect_pdf(built["PDF"]), inspect_pdf(data)
    assert not after["errors"] and after["page_count"] == before["page_count"] + 1
    assert after["outline_items"] == before["outline_items"] and after["links"] == before["links"]
    assert after["media_boxes"][0] == before["media_boxes"][0]
    reader = pypdf.PdfReader(io.BytesIO(data))
    assert "THE SALT" in reader.pages[0].extract_text() and "MARA VELL" in reader.pages[0].extract_text()
    assert "The salt was white" not in reader.pages[0].extract_text()
    assert "The salt was white" in "".join(pg.extract_text() for pg in reader.pages[1:])
    # outline destinations still land on the same (now shifted) pages
    first = reader.outline[0]
    assert reader.get_destination_page_number(first) >= 1
    assert (b"/Subtype /Image" in data) == with_art


@pytest.mark.parametrize("with_art", [True, False])
def test_epub_export_gets_a_cover(exported, with_art):
    app, pid, _built = exported
    if with_art:
        app.cover.upload("author", pid, BR, png(300, 450, fmt="JPEG"))
    data, mime, _name = cover.book_with_cover(app, pid, BR, "EPUB")
    assert mime == "application/epub+zip"
    rep = inspect_epub(data)
    assert not rep["errors"], rep["errors"]
    z = zipfile.ZipFile(io.BytesIO(data))
    assert z.namelist()[0] == "mimetype"
    opf = z.read("OEBPS/content.opf").decode()
    assert "<meta name=\"cover\" content=\"cover-image\"/>" in opf and "properties=\"cover-image\"" in opf
    assert re.search(r"<spine[^>]*>\s*<itemref idref=\"cover\"/>", opf)
    assert "href=\"text/cover.xhtml\" media-type=\"application/xhtml+xml\" properties=\"svg\"" in opf
    xhtml = z.read("OEBPS/text/cover.xhtml").decode()
    assert "epub:type=\"cover\"" in xhtml and ">THE SALT<" in xhtml and "<svg" in xhtml
    import lxml.etree as LET
    LET.fromstring(xhtml.encode())  # well-formed XHTML
    if with_art:
        assert "OEBPS/images/cover-art.jpg" in z.namelist() and "../images/cover-art.jpg" in xhtml
    else:
        assert "OEBPS/images/cover.svg" in z.namelist() and "image/svg+xml" in opf
    # a second pass leaves it alone
    assert cx.add_cover_epub(data, svg_inline="<svg/>", art=None, svg_file="<svg/>", title="x") == data


@pytest.mark.parametrize("with_art", [True, False])
def test_docx_export_gets_a_cover_page(exported, with_art):
    app, pid, _built = exported
    cover.save(app, pid, BR, {"texts": {"subtitle": "A desert crossing", "series": "Salt Books"}})
    if with_art:
        app.cover.upload("author", pid, BR, png(300, 450))
    data, mime, _name = cover.book_with_cover(app, pid, BR, "DOCX")
    assert mime.endswith("wordprocessingml.document")
    assert not inspect_docx(data)["errors"]
    doc = python_docx.Document(io.BytesIO(data))
    texts = [p.text for p in doc.paragraphs[:8]]
    assert texts[0] == "Salt Books" and texts[1] == "The Salt Road" and texts[2] == "A desert crossing"
    assert "Mara Vell" in texts
    assert doc.paragraphs[1].style.name == "Title"
    assert (len(doc.inline_shapes) >= 1) == with_art
    assert any('w:type="page"' in p._p.xml for p in doc.paragraphs[:8])


def test_book_with_cover_needs_an_export(fx):
    app, pid = fx
    with pytest.raises(AkoweError) as e:
        cover.book_with_cover(app, pid, BR, "EPUB")
    assert "export the book first" in str(e.value)
    with pytest.raises(AkoweError):
        cover.book_with_cover(app, pid, BR, "MOBI")


# ====================================================================== book save
def test_book_save_includes_the_cover(exported):
    app, pid, _built = exported
    files, _info = app.saver.collect(pid, BR)
    assert not [k for k in files if k.startswith("Cover/")]  # never opened on the cover page
    cover.save(app, pid, BR, {"template": "fantasy", "print": {"page_count": 240}})
    app.cover.upload("author", pid, BR, png(300, 450))
    files, _info = app.saver.collect(pid, BR)
    cov = sorted(k for k in files if k.startswith("Cover/"))
    assert {"Cover/cover-design.json", "Cover/ebook cover.pdf", "Cover/ebook cover.svg", "Cover/cover art.png",
            "Cover/print cover.pdf"} <= set(cov)
    assert {"Cover/The Salt Road (with cover).pdf", "Cover/The Salt Road (with cover).epub",
            "Cover/The Salt Road (with cover).docx"} <= set(cov)
    assert all(len(k) <= 60 for k in cov), cov
    assert json.loads(files["Cover/cover-design.json"])["template"] == "fantasy"
    assert inspect_pdf(files["Cover/print cover.pdf"])["page_count"] == 1


# ====================================================================== HTTP
def test_http_api_contract(fx, no_provider_sleep):  # noqa: F811
    from akowe.server import make_handler
    app, pid = fx
    b = FakeImageBridge()
    forge.save_bridge_settings(app, {"url": b.url, "token": "t"})
    make_char(app, pid, "Mara Vell")
    httpd = ThreadingHTTPServer(("127.0.0.1", 0), make_handler(app))
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    port = httpd.server_address[1]
    base = f"/api/projects/{pid}/branches/{BR}/cover"
    try:
        st, d, _ = call(port, "GET", base)
        assert st == 200 and {"design", "texts", "templates", "trims", "papers", "print", "characters"} <= set(d)
        assert d["characters"] == ["Mara Vell"]
        st, d, _ = call(port, "PUT", base, {"template": "romance", "print": {"page_count": 180}, "revision": 0})
        assert st == 200 and d["design"]["revision"] == 1 and d["print"]["ok"]
        st, e, _ = call(port, "PUT", base, {"template": "romance", "revision": 0})
        assert st == 409
        st, e, _ = call(port, "PUT", base, {"template": "nope"})
        assert st == 400 and e["errors"][0]["code"] == "INVALID_INPUT" and "template" in e["errors"][0]["message"]
        st, j, _ = call(port, "POST", f"{base}/art", {"engine": "chatgpt"})
        assert st == 202 and j["job_id"]
        assert wait(app, j)["status"] == "done" and b.requests[-1]["purpose"] == "cover"
        st, j2, _ = call(port, "GET", f"{base}/jobs/{j['job_id']}")
        assert st == 200 and j2["status"] == "done"
        st, u, _ = call(port, "PUT", f"{base}/art", png(), content_type="image/png")
        assert st == 200 and u["design"]["art"]["source"] == "upload"
        st, pdf, resp = call(port, "GET", f"{base}/ebook?download=1")
        assert st == 200 and pdf.startswith(b"%PDF") and "attachment" in resp.getheader("Content-Disposition")
        st, svg, resp = call(port, "GET", f"{base}/ebook?format=svg")
        assert st == 200 and resp.getheader("Content-Type").startswith("image/svg+xml") and b"<svg" in svg
        st, art, _ = call(port, "GET", f"{base}/ebook?format=art")
        assert st == 200 and art == png()
        st, pdf, resp = call(port, "GET", f"{base}/print?trim=6x9&paper=cream")
        assert st == 200 and pdf.startswith(b"%PDF")
        assert inspect_pdf(pdf)["media_boxes"][0][2] == pytest.approx((12.25 + 180 * 0.0025) * 72, abs=0.01)
        st, e, _ = call(port, "GET", f"{base}/print?trim=6x9&paper=cream&pages=10")
        assert st == 400 and "at least 24" in e["errors"][0]["message"]
        st, e, _ = call(port, "GET", f"{base}/print?trim=8x10")
        assert st == 400 and "trim size" in e["errors"][0]["message"]
        st, svg, _ = call(port, "POST", f"{base}/preview?what=print", {"template": "scifi", "blurb": "Unsaved blurb."})
        assert st == 200 and b"Unsaved blurb." in svg and b"class=\"guide fold\"" in svg
        st, svg, _ = call(port, "POST", f"{base}/preview", {"texts": {"title": "Preview Title"}})
        assert st == 200 and b"Preview Title" in svg.replace(b"PREVIEW TITLE", b"Preview Title")
        assert cover.get_design(app.store, pid, BR)["template"] == "romance"  # a preview saves nothing
        assert "blurb" not in get_profile(app.store, pid)
        st, e, _ = call(port, "GET", f"{base}/book?format=EPUB")
        assert st == 400 and "export the book first" in e["errors"][0]["message"]
        st, e, _ = call(port, "GET", f"{base}/nope")
        assert st == 404
        st, _b, _ = call(port, "GET", base, token="wrong-token")
        assert st == 401
    finally:
        httpd.shutdown()
        httpd.server_close()
        b.close()


def test_rendered_png_samples(tmp_path):
    """The rendered covers really are pictures (and the docs' samples come from the same code)."""
    art = load_image(png(400, 600, (90, 60, 40)))
    pdf, _o, _i = render_front({"template": "literary", "texts": TEXTS}, art)
    img = pdfium.PdfDocument(pdf)[0].render(scale=0.4).to_pil()
    assert abs(img.size[0] - cx.EBOOK_PT[0] * 0.4) <= 1 and abs(img.size[1] - cx.EBOOK_PT[1] * 0.4) <= 1
    assert isinstance(img, Image.Image)
