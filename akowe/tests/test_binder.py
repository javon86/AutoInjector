"""📖 Book Bible binder (data side): the mapping table covers every schema field id, the contract's shape on a
demo book (a Make-a-book run with the fake AI + the Moki Vey example world + a faction), the mapping and the auto
fields, the looks, the 🆕 fields (storage, validation, migration, journal), "✨ Fill from the book" (proposals that
never overwrite), the Word and PDF exports (a fake AutoInjector ``POST /render/pdf``), the book save and the HTTP
API."""
from __future__ import annotations

import base64
import io
import json
import threading
import time
import types
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import pytest

from akowe import binder, binder_fields as bf, binder_html, forge, forge_example, journal
from akowe import providers as prov_mod
from akowe.book_path import BookPath, get_profile, save_profile
from akowe.errors import AkoweError

from autobook_fake import BookFakeAI
from test_autobook import make_app

BR = "main"
STATUSES = {"approved", "proposed", "supplied", "empty", "auto"}
PDF = b"%PDF-1.7\n% fake binder pdf\n%%EOF\n"


# ====================================================================== helpers / fixtures
def _make_book(tmp):
    old = prov_mod.time
    prov_mod.time = types.SimpleNamespace(sleep=lambda s: None, monotonic=time.monotonic, time=time.time)
    try:
        app = make_app(tmp, fake=BookFakeAI())
        r = app.autobook.start("author", {"author_confirms": True})
        app.autobook.wait_for(r["id"], 240)
        run = app.autobook.get(r["id"])
        assert run["status"] == "done", run.get("problem")
    finally:
        prov_mod.time = old
    pid = run["project_id"]
    forge.save_bridge_settings(app, {"url": "http://127.0.0.1:9"})  # nothing listens there: no AutoInjector
    forge_example.load(app, "author", pid, BR, {})                  # Moki Vey + the item, place and magic system
    bp = BookPath(app.engine, "author", pid, BR)
    cand = {"concept": "The Harbour Board", "attributes": {"membership": "Twelve captains and the harbourmaster",
                                                           "hierarchy": "The harbourmaster chairs it",
                                                           "goals": "Keep the harbour open at any cost",
                                                           "alliances": "The shipping guild",
                                                           "rivalries": "The lighthouse keepers"},
            "linked_refs": [], "protected_fact_flags": [],
            "facts": [{"subject_id": "harbour-board", "predicate": "entity.name", "value": "The Harbour Board",
                       "assertion_type": "FACT", "effective_from": None, "effective_to": None, "source_refs": []}]}
    bp.create("A06", {"concept": "The Harbour Board"}, manual_candidate=cand)
    return app, pid


@pytest.fixture(scope="module")
def demo(tmp_path_factory):
    app, pid = _make_book(tmp_path_factory.mktemp("binder_demo"))
    return app, pid, binder.build(app, pid, BR)


@pytest.fixture
def book(tmp_path):
    return _make_book(tmp_path)


def page_of(data, title=None, template=None):
    return next(p for p in data["pages"] if (title is None or p["title"] == title)
                and (template is None or p["template"] == template))


def fld(page, fid):
    code = fid.rsplit(".", 1)[0]
    return page["sections"][code]["fields"][fid]


def char_rid(app, pid, name):
    from akowe import characters as chars_mod
    return next(c["record_id"] for c in chars_mod.canon_characters(app.store, pid, BR) if c["name"] == name)


# ====================================================================== the mapping table
def test_every_schema_id_is_mapped_or_deliberately_unavailable():
    ids = binder.all_ids()
    assert len(ids) == len(set(ids)) == 480
    assert set(binder.MAP) == set(ids), (set(ids) ^ set(binder.MAP))
    for i, spec in binder.MAP.items():
        if spec is None:
            assert binder.UNAVAILABLE.get(i), f"{i} is None without a reason"
            continue
        for s in spec if isinstance(spec, tuple) else (spec,):
            kind, _, key = s.partition(":")
            assert kind in ("profile", "new", "cover", "choices", "plan", "forge", "place", "magic", "item",
                            "canon", "pic", "row", "auto", "const"), s
            if kind == "auto":
                assert key in binder.AUTO, s
            if kind == "new":
                assert any(key in rules for rules in bf.FIELDS.values()), s
            if kind == "forge":
                assert key in forge.FIELDS, s
            if kind in ("place", "magic", "item"):
                from akowe.forge_kinds_spec import SPECS
                assert key in SPECS[kind]["fields"], s
    assert set(binder.UNAVAILABLE) == {i for i, s in binder.MAP.items() if s is None}
    # row specs exactly in the repeating sections, and every repeating section has a row producer
    rep = {s["code"] for t in binder.schema().values() for s in t["sections"] if s.get("repeat")}
    assert set(binder.ROWS) == rep and all(callable(binder.ROW_FUNCS[n]) for n in binder.ROWS.values())
    for i, spec in binder.MAP.items():
        is_row = isinstance(spec, str) and spec.startswith("row:")
        assert is_row == (i.rsplit(".", 1)[0] in rep), i
    # every 🆕 field of the spec is stored by Akọwe (marked new) — none is left unavailable
    new_in_spec = {f["id"] for t in binder.schema().values() for s in t["sections"] for f in s["fields"] if f["new"]}
    assert new_in_spec <= binder.NEW_IDS
    kinds = binder.source_kinds()
    assert set(kinds) == set(ids) and kinds["5.2.5"] == "new" and kinds["5.4.2"] == "none"


def test_look_default_per_genre():
    want = {"Fantasy": "fantasy", "Science fiction": "scifi", "Mystery": "noir", "Thriller": "noir",
            "Romance": "romance", "Horror": "horror", "Historical": "classic", "Literary": "classic",
            "Adventure": "fantasy", "Comedy": "playful", "Children’s picture book": "playful"}
    for genre, look in want.items():
        assert binder.look_for_genre(genre) == look, genre
    assert binder.look_for_genre("Cozy mystery") == "noir" and binder.look_for_genre("space opera") == "scifi"
    assert binder.look_for_genre("") == binder.look_for_genre("Westerns") == "classic"
    assert set(binder.LOOKS) == set(binder_html.LOOK_VARS) and len(binder.LOOKS) == 7


# ====================================================================== the contract
def test_contract_shape_every_template_and_id(demo):
    app, pid, d = demo
    assert d["schema_version"] == 1 and d["book"]["title"] == get_profile(app.store, pid)["title"]
    assert d["book"]["look"] == d["book"]["look_default"] == "noir" and d["book"]["look_stored"] is None
    assert {p["template"] for p in d["pages"]} == set(binder.schema())
    assert [p["page"] for p in d["pages"]] == list(range(1, len(d["pages"]) + 1))
    assert len({p["page_id"] for p in d["pages"]}) == len(d["pages"])
    for p in d["pages"]:
        tpl = binder.schema()[p["template"]]
        assert set(p) >= {"template", "entry_id", "title", "page_id", "status", "sections", "pics", "page"}
        assert p["status"] in ("approved", "draft", "needs_you")
        assert set(p["sections"]) == {s["code"] for s in tpl["sections"]}
        for sec in tpl["sections"]:
            s = p["sections"][sec["code"]]
            assert set(s["fields"]) == {f["id"] for f in sec["fields"]}, (p["title"], sec["code"])
            for f in s["fields"].values():
                assert f["status"] in STATUSES and "value" in f
            if sec.get("repeat"):
                assert isinstance(s["rows"], list)
                for row in s["rows"]:
                    assert {k for k in row if not k.startswith("_")} == {f["id"] for f in sec["fields"]}
                    assert all(v["status"] in STATUSES for k, v in row.items() if not k.startswith("_"))
            else:
                assert "rows" not in s
        for k, pic in p["pics"].items():
            assert pic["url"].startswith(f"/api/projects/{pid}/branches/{BR}/assets/")
    first = {c["template"]: c["page"] for c in d["contents"]}
    assert first["1"] == 1 and first["5"] == page_of(d, template="5")["page"]
    json.dumps(d)  # plain JSON


def test_mapping_from_the_demo_book(demo):
    app, pid, d = demo
    prof = get_profile(app.store, pid)
    cover = page_of(d, template="1")
    assert fld(cover, "1.1.1")["value"] == prof["title"] and fld(cover, "1.1.1")["source"] == "profile:title"
    assert fld(cover, "1.3.1")["value"] == "BOOK BIBLE" and fld(cover, "1.3.5")["value"] == "3 chapters"
    glance = page_of(d, template="3")
    b01 = BookPath(app.engine, "author", pid, BR).approved("B01")[0]["data"]
    assert fld(glance, "3.3.1")["value"] == b01["premise"] and fld(glance, "3.3.1")["status"] == "approved"
    assert fld(glance, "3.7.5")["value"] == b01["dramatic_question"]  # no central question yet: the B01 question
    moki = page_of(d, "Moki Vey")
    dossier = forge.public(app.store, pid, BR, char_rid(app, pid, "Moki Vey"))
    phys = next(f for s in dossier["sections"] for f in s["fields"] if f["key"] == "physique")
    assert fld(moki, "5.6.2")["value"] == phys["text"] and fld(moki, "5.6.2")["status"] == "proposed"
    assert fld(moki, "5.6.2")["source"] == "forge:physique"
    assert fld(moki, "5.6.2")["edit"]["route"].endswith(f"/forge/{char_rid(app, pid, 'Moki Vey')}")
    assert moki["status"] == "needs_you" and "5.1.1" in moki["pics"] and "5.3.2" in moki["pics"]
    assert fld(moki, "5.1.4")["value"].startswith(dossier["code"])
    assert moki["sections"]["5.8"]["rows"][0]["5.8.2"]["value"].startswith("Imani Vey")
    assert fld(moki, "5.4.2")["unavailable"] and fld(moki, "5.4.2")["status"] == "empty"
    mara = page_of(d, "Mara Vell")
    assert fld(mara, "5.4.3")["value"].startswith("Mara Vell wants")  # the character sheet, through the dossier
    assert fld(mara, "5.7.11")["value"] == "Who signed the log?"      # voice card sample line 1
    nacre = page_of(d, "Nacre")
    assert fld(nacre, "7.1.3")["value"] == "PL-001" and "7.1.1" in nacre["pics"]
    harbor = page_of(d, "Gullhaven Harbor")  # an A04 record without a dossier: the canon attributes
    assert fld(harbor, "7.2.1")["source"] == "canon:A04.geography" and fld(harbor, "7.2.1")["status"] == "approved"
    assert fld(page_of(d, template="8A"), "8A.1.3")["value"] == "MG-001"
    assert fld(page_of(d, template="8B"), "8B.1.3")["value"] == "AR-001"
    board = page_of(d, "The Harbour Board")
    assert fld(board, "8C.5.1")["value"] == "Twelve captains and the harbourmaster"
    assert fld(board, "8C.5.3")["value"] == "The lighthouse keepers" and board["status"] == "draft"
    assert fld(board, "8C.1.1")["unavailable"]
    ch1 = page_of(d, template="11")
    assert fld(ch1, "11.2.1")["value"] == "Mara Vell" and fld(ch1, "11.1.4")["value"] == "Approved"
    assert ch1["sections"]["11.4"]["rows"][0]["11.4.3"]["value"].startswith("The evening tide")
    rules = page_of(d, template="12")
    assert "brass key" in fld(rules, "12.2.1")["value"] and rules["sections"]["12.5"]["rows"]
    assert any("Supplied canon" in str(r["12.1.2"]["value"]) for r in rules["sections"]["12.1"]["rows"])


def test_auto_fields(demo):
    app, pid, d = demo
    mara = page_of(d, "Mara Vell")
    assert fld(mara, "5.14.1")["value"] == "Chapter 1" and fld(mara, "5.14.2")["value"] == "Ch 1, Ch 2, Ch 3"
    assert fld(mara, "5.14.1")["status"] == "auto"
    cast = page_of(d, template="4")
    codes = [r["4.2.3"]["value"] for r in cast["sections"]["4.2"]["rows"]]
    assert sorted(codes) == ["CH-001", "CH-002", "CH-003"] and fld(cast, "4.1.2")["value"] == "3 characters"
    cover = page_of(d, template="1")
    pct = fld(cover, "1.4.1")["value"]
    assert pct.endswith("%") and 0 < int(pct[:-1]) < 100 and "missing" in fld(cover, "1.4.3")["value"]
    assert fld(cover, "1.5.2")["value"] == "Book Bible schema 1"
    mp = page_of(d, template="9")
    harbor = page_of(d, "Gullhaven Harbor")
    key = next(r for r in mp["sections"]["9.2"]["rows"] if r["9.2.2"]["value"] == "Gullhaven Harbor")
    assert key["9.2.3"]["value"] == str(harbor["page"])
    contents = page_of(d, template="2")
    rows = {r["2.1.1"]["value"].split(" ", 1)[1]: r for r in contents["sections"]["2.1"]["rows"]}
    assert len(rows) == 9 and rows["Chapters"]["2.1.3"]["value"] == str(page_of(d, template="11")["page"])
    assert rows["Cast"]["2.1.2"]["value"] == "3 characters"
    todo = contents["sections"]["2.3"]["rows"]
    assert todo and todo[0]["2.3.1"]["value"].startswith("Moki Vey") and todo[0]["2.3.3"]["value"] == "✔ Approve"
    assert any(t["action"] == "make_picture" for t in d["todo"])
    plan = page_of(d, template="10")
    assert fld(plan, "10.2.1")["value"].startswith("Chapter 1") and fld(plan, "10.2.8")["value"].startswith("Chapter 3")
    rel = page_of(d, template="6")
    assert rel["graph"]["centre"] == char_rid(app, pid, "Mara Vell") and rel["graph"]["edges"]
    assert fld(rel, "6.1.3")["value"] == "Mara Vell"


# ====================================================================== 🆕 fields
def test_new_fields_storage_validation_and_journal(book):
    app, pid = book
    moki = char_rid(app, pid, "Moki Vey")
    t = f"character:{moki}"
    out = bf.edit(app.store, pid, BR, {"target": t, "fields": {"height": "1.68 m", "stat_strength": 7,
                                                               "handedness": "right"}}, app=app)
    assert out["fields"]["height"]["value"] == "1.68 m" and out["fields"]["stat_strength"]["value"] == 7
    with pytest.raises(AkoweError) as e:
        bf.edit(app.store, pid, BR, {"target": t, "fields": {"stat_speed": 11}}, app=app)
    assert e.value.code == "INVALID_INPUT" and "1 to 10" in str(e.value.errors[0]["message"])
    with pytest.raises(AkoweError):
        bf.edit(app.store, pid, BR, {"target": t, "fields": {"height": "x" * 61}}, app=app)
    with pytest.raises(AkoweError) as e:
        bf.edit(app.store, pid, BR, {"target": t, "fields": {"height": "2 m"}, "version": 0}, app=app)
    assert e.value.code == "VERSION_CONFLICT"
    with pytest.raises(AkoweError) as e:
        bf.edit(app.store, pid, BR, {"target": "character:A01-nobody", "fields": {"height": "2 m"}}, app=app)
    assert e.value.code == "NOT_FOUND"
    data = binder.build(app, pid, BR)
    harbor = page_of(data, "Gullhaven Harbor")
    with pytest.raises(AkoweError):
        bf.edit(app.store, pid, BR, {"target": harbor["target"], "fields": {"danger_level": 6}}, app=app)
    bf.edit(app.store, pid, BR, {"target": harbor["target"], "fields": {"danger_level": "3"}}, app=app)
    board = page_of(data, "The Harbour Board")
    bf.edit(app.store, pid, BR, {"target": board["target"], "fields": {"symbol_colours": "A gull on navy"}}, app=app)
    ch1 = page_of(data, template="11")
    bf.edit(app.store, pid, BR, {"target": ch1["target"], "fields": {"time_since_previous": "two days later"}},
            app=app)
    rel = bf.rel_target(moki, "Imani Vey")
    bf.edit(app.store, pid, BR, {"target": rel, "fields": {"type": "Family"}}, app=app)
    with pytest.raises(AkoweError):
        bf.edit(app.store, pid, BR, {"target": rel, "fields": {"type": "frenemy"}}, app=app)
    # the book's 🆕 values live in the book profile
    bf.edit(app.store, pid, BR, {"target": "book", "fields": {"series": "Harbour Lights", "series_number": 2,
                                                              "themes": ["duty", "grief"],
                                                              "central_question": "Who keeps the light?"}}, app=app)
    prof = get_profile(app.store, pid)
    assert prof["series"] == "Harbour Lights" and prof["themes"] == ["duty", "grief"] and prof["series_number"] == 2
    with pytest.raises(AkoweError):
        save_profile(app.store, pid, {"themes": ["a", "b", "c", "d"]})
    with pytest.raises(AkoweError):
        bf.edit(app.store, pid, BR, {"target": "book", "fields": {"series_number": 0}}, app=app)
    d2 = binder.build(app, pid, BR)
    m = page_of(d2, "Moki Vey")
    assert fld(m, "5.2.5") == dict(fld(m, "5.2.5"), value="1.68 m", status="supplied", new=True)
    assert fld(m, "5.2.5")["edit"]["target"] == t and fld(m, "5.9.7")["value"] == 7
    assert m["sections"]["5.8"]["rows"][0]["5.8.3"]["value"] == "family"
    assert fld(page_of(d2, "Gullhaven Harbor"), "7.2.7")["value"] == 3
    assert fld(page_of(d2, "The Harbour Board"), "8C.3.5")["value"] == "A gull on navy"
    assert fld(page_of(d2, template="11"), "11.2.4")["value"] == "two days later"
    cover, glance = page_of(d2, template="1"), page_of(d2, template="3")
    assert fld(cover, "1.1.4")["value"] == "Harbour Lights, book 2"
    assert fld(glance, "3.7.2")["value"] == "duty" and fld(glance, "3.7.5")["value"] == "Who keeps the light?"
    assert page_of(d2, template="6")["graph"]["edges"]  # the typed tie is drawn (Imani is not a character: not here)
    entries = journal.build(app, pid, BR)["entries"]
    assert sum(1 for e in entries if e.get("kind") == "binder") >= 6


def test_migration_keeps_old_books_working(book):
    app, pid = book
    moki = char_rid(app, pid, "Moki Vey")
    # the first layout: plain values, no status, one value that no longer validates, unknown keys and targets
    app.store.set_setting(f"bible_fields:{pid}:{BR}", {"targets": {
        f"character:{moki}": {"height": "1.7 m", "stat_strength": 99, "favourite_food": "soup"},
        "spaceship:x": {"height": "1"}, "chapter:ch1": "garbage"}})
    doc = bf.get_doc(app.store, pid, BR)
    assert doc["schema"] == bf.SCHEMA and list(doc["targets"]) == [f"character:{moki}"]
    assert doc["targets"][f"character:{moki}"]["fields"] == {"height": {"value": "1.7 m", "status": "supplied",
                                                                        "protected": False}}
    assert fld(page_of(binder.build(app, pid, BR), "Moki Vey"), "5.2.5")["value"] == "1.7 m"
    # an edit writes the current schema back; garbage and no document read as empty
    bf.edit(app.store, pid, BR, {"target": f"character:{moki}", "fields": {"weight": "65 kg"}}, app=app)
    assert app.store.get_setting(f"bible_fields:{pid}:{BR}")["schema"] == bf.SCHEMA
    assert bf.migrate("nonsense") == bf.migrate(None) == {"schema": bf.SCHEMA, "revision": 0, "targets": {}}
    # a profile from before the binder (no new keys) and an unknown stored look still build
    prof = app.store.get_setting(f"book_profile:{pid}")
    prof["binder_look"] = "vaporwave"
    app.store.set_setting(f"book_profile:{pid}", prof)
    assert binder.build(app, pid, BR)["book"]["look"] == "noir"


# ====================================================================== ✨ Fill from the book
def test_fill_from_the_book_proposes_and_never_overwrites(book, monkeypatch):
    app, pid = book
    moki = char_rid(app, pid, "Moki Vey")
    t = f"character:{moki}"
    forge.approve(app.store, pid, BR, moki, {"fields": ["physique"]})   # now it is the book's own text
    bf.edit(app.store, pid, BR, {"target": t, "fields": {"height": "1.70 m"}}, app=app)        # the author's
    bf.edit(app.store, pid, BR, {"target": t, "protected": {"build": True}}, app=app)         # protected, empty
    bf.edit(app.store, pid, BR, {"target": "book", "fields": {"subgenre": "Cosy mystery"}}, app=app)
    asked = []

    def fake_ai(job, *, op_for_roles, task, schema, context, protected):
        asked.append((context["subject"], sorted(schema["properties"])))
        assert "Book Bible quick facts" in task
        return {k: ("7" if k.startswith("stat_") else "Invented" if k != "subgenre" else "Space opera")
                for k in schema["properties"]} | ({"rel_1": "family"} if "rel_1" in schema["properties"] else {})
    monkeypatch.setattr(app.assist, "_call_ai", fake_ai)
    j = bf.start_fill(app, "author", pid, BR, {"targets": [t, "book"]})
    j = forge.wait(app, "author", j["job_id"])
    assert j["status"] == "done", j
    vals = bf.values(app.store, pid, BR, t)
    assert vals["height"]["value"] == "1.70 m" and vals["height"]["status"] == "supplied"        # kept
    assert vals["build"]["value"] is None and vals["build"]["protected"]                          # never touched
    assert vals["weight"] == dict(vals["weight"], value="65 kg (143 lb)", status="proposed", basis="physique")
    assert vals["stat_strength"]["value"] == 7 and vals["stat_strength"]["status"] == "proposed"
    subjects = dict(asked)
    assert "height" not in subjects["Moki Vey"] and "build" not in subjects["Moki Vey"]
    assert "weight" not in subjects["Moki Vey"]  # Akọwe's own reading of Physique came first
    assert get_profile(app.store, pid)["subgenre"] == "Cosy mystery"
    assert "subgenre" not in subjects[get_profile(app.store, pid)["title"]]
    rel = bf.values(app.store, pid, BR, bf.rel_target(moki, "Imani Vey"))["type"]
    assert rel["value"] == "family" and rel["status"] == "proposed"
    # a second run proposes nothing over the waiting proposals
    j2 = forge.wait(app, "author", bf.start_fill(app, "author", pid, BR, {"targets": [t], "ai": False})["job_id"])
    assert j2["status"] == "done" and bf.values(app.store, pid, BR, t)["weight"]["value"] == "65 kg (143 lb)"
    # the binder shows them as proposals; approving makes them approved; book proposals go to the profile
    page = page_of(binder.build(app, pid, BR), "Moki Vey")
    assert fld(page, "5.2.6")["status"] == "proposed"
    out = bf.approve(app.store, pid, BR, {"target": t, "fields": ["weight"]})
    assert out["fields"]["weight"]["status"] == "approved"
    assert bf.propose(app.store, pid, BR, "book", {"central_question": "Will the light return?", "subgenre": "x"}) \
        == []  # the AI's central question waits already; the sub-genre is the author's
    assert bf.values(app.store, pid, BR, "book")["central_question"]["status"] == "proposed"
    bf.approve(app.store, pid, BR, {"target": "book", "all": True})
    prof = get_profile(app.store, pid)
    assert prof["central_question"] == "Invented" and prof["subgenre"] == "Cosy mystery"
    assert bf.values(app.store, pid, BR, "book")["central_question"]["status"] == "approved"
    assert bf.dismiss(app.store, pid, BR, {"target": t, "all": True})["dismissed"]
    assert bf.values(app.store, pid, BR, t)["stat_strength"]["value"] is None


# ====================================================================== exports
class FakePdfBridge:
    """AutoInjector's POST /render/pdf (and nothing else: /render/image etc. answer 404)."""

    def __init__(self, status=200):
        self.calls, self.status = [], status
        outer = self

        class H(BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def do_POST(self):
                body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
                outer.calls.append((self.path, body, self.headers.get("Authorization")))
                if self.path != "/render/pdf":
                    return self._send(404, {"error": "NOT_FOUND"})
                if outer.status != 200:
                    return self._send(outer.status, {"ok": False, "error": "RENDER_NOT_WIRED", "message": "no"})
                self._send(200, {"ok": True, "mime": "application/pdf", "pdfBase64": base64.b64encode(PDF).decode(),
                                 "bytes": len(PDF), "pages": 3})

            def _send(self, code, obj):
                raw = json.dumps(obj).encode()
                self.send_response(code)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(raw)))
                self.end_headers()
                self.wfile.write(raw)
        self.srv = ThreadingHTTPServer(("127.0.0.1", 0), H)
        self.url = f"http://127.0.0.1:{self.srv.server_address[1]}"
        threading.Thread(target=self.srv.serve_forever, daemon=True).start()

    def close(self):
        self.srv.shutdown()


def test_docx_export_opens_with_sections_pictures_and_look_colour(demo):
    import docx as python_docx
    app, pid, d = demo
    data, name = binder.export_docx(app, pid, BR)
    assert name.endswith("- Book Bible.docx")
    doc = python_docx.Document(io.BytesIO(data))
    h1 = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert len(h1) == len(d["pages"]) and any("Moki Vey" in x for x in h1) and any("Rules & facts" in x for x in h1)
    h2 = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert any(x.startswith("📋 Quick stats") for x in h2) and any(x.startswith("🗣 Voice card") for x in h2)
    assert len(doc.tables) > 100 and len(doc.inline_shapes) >= 10  # Moki's portrait and views, the world's art
    cells = " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "Height" in cells and "Gullhaven Harbor" in cells
    accent = binder_html.LOOK_VARS["noir"][2].lstrip("#").upper()
    xml = doc.styles["Heading 1"].element.xml
    assert accent in xml
    other, _ = binder.export_docx(app, pid, BR, look="fantasy")
    assert binder_html.LOOK_VARS["fantasy"][2].lstrip("#").upper().encode() in other


def test_pdf_without_autoinjector_and_with_a_fake_bridge(demo):
    app, pid, _d = demo
    with pytest.raises(AkoweError) as e:
        binder.export_pdf(app, pid, BR)
    assert e.value.errors[0]["problem"] == "RENDER_NOT_WIRED"
    assert "Open AutoInjector to make the PDF, or use 🖨 Print → Save as PDF" in e.value.errors[0]["message"]
    br = FakePdfBridge()
    try:
        forge.save_bridge_settings(app, {"url": br.url, "token": "pdf-token"})
        pdf, name = binder.export_pdf(app, pid, BR, page_size="Letter")
        assert pdf == PDF and name.endswith("- Book Bible.pdf")
        path, body, auth = br.calls[-1]
        assert path == "/render/pdf" and auth == "Bearer pdf-token"
        assert body["pageSize"] == "Letter" and body["printBackground"] is True and "Book Bible" in body["title"]
        html = body["html"]
        assert html.startswith("<!doctype html>") and "<script" not in html and 'src="/api/' not in html
        assert html.count('src="data:image/') >= 10 and "--accent:" in html and 'data-id="5.2.5"' in html
    finally:
        forge.save_bridge_settings(app, {"url": "http://127.0.0.1:9"})
        br.close()
    br = FakePdfBridge(status=501)  # an AutoInjector without the route
    try:
        forge.save_bridge_settings(app, {"url": br.url})
        with pytest.raises(AkoweError) as e:
            binder.export_pdf(app, pid, BR)
        assert e.value.errors[0]["problem"] == "RENDER_NOT_WIRED"
    finally:
        forge.save_bridge_settings(app, {"url": "http://127.0.0.1:9"})
        br.close()


def test_html_renderer_uses_each_look():
    data = {"book": {"title": "T", "look": "horror"}, "pages": [], "generated_at": "x"}
    html = binder_html.render(data)
    assert binder_html.LOOK_VARS["horror"][2] in html and 'data-look="horror"' in html
    assert "@page{size:A4" in html and "@page{size:Letter" in binder_html.render(data, page_size="Letter")


# ====================================================================== the book save
def test_book_save_writes_the_binder_files(book, tmp_path):
    app, pid = book
    app.saver.save_settings({"path": str(tmp_path / "books")})
    res = app.saver.save_now(pid, BR)
    folder = tmp_path / "books" / res["folder"]
    assert (folder / "Book Bible" / "Book Bible.docx").read_bytes()[:2] == b"PK"
    assert not (folder / "Book Bible" / "Book Bible.pdf").exists()
    assert res["not_made"] and res["not_made"][0]["file"] == "Book Bible/Book Bible.pdf"
    assert "Open AutoInjector" in res["not_made"][0]["why"]
    br = FakePdfBridge()
    try:
        forge.save_bridge_settings(app, {"url": br.url})
        res = app.saver.save_now(pid, BR)
        assert (folder / "Book Bible" / "Book Bible.pdf").read_bytes() == PDF and not res["not_made"]
        n = sum(1 for c in br.calls if c[0] == "/render/pdf")
        app.saver.save_now(pid, BR)  # unchanged binder: the kept PDF, no new render
        assert sum(1 for c in br.calls if c[0] == "/render/pdf") == n
    finally:
        br.close()
    # AutoInjector gone again: the PDF saved before stays, the DOCX is still written
    forge.save_bridge_settings(app, {"url": "http://127.0.0.1:9"})
    bf.edit(app.store, pid, BR, {"target": "book", "fields": {"subgenre": "Harbour noir"}}, app=app)
    res = app.saver.save_now(pid, BR)
    assert (folder / "Book Bible" / "Book Bible.pdf").read_bytes() == PDF
    assert res["not_made"] and (folder / "Book Bible" / "Book Bible.docx").exists()


# ====================================================================== HTTP
def test_http_api(book):
    app, pid = book

    def call(method, path, body=None, query=None):
        return app.route(method, f"/api/projects/{pid}/branches/{BR}/binder{path}", query or {}, body, "author", {})
    st, d = call("GET", "")
    assert st == 200 and d["book"]["look"] == "noir"
    st, d = call("GET", "", query={"look": ["fantasy"]})
    assert d["book"]["look"] == "fantasy" and d["book"]["look_stored"] is None
    st, out = call("PUT", "/look", {"look": "romance"})
    assert out == {"ok": True, "look": "romance", "look_stored": "romance", "look_default": "noir"}
    assert call("GET", "")[1]["book"]["look"] == "romance"
    assert call("PUT", "/look", {"look": None})[1]["look"] == "noir"
    with pytest.raises(AkoweError):
        call("PUT", "/look", {"look": "vaporwave"})
    with pytest.raises(AkoweError):
        call("GET", "", query={"look": ["vaporwave"]})
    st, sch = call("GET", "/schema")
    assert set(sch["templates"]) == set(binder.schema()) and sch["unavailable"]
    moki = char_rid(app, pid, "Moki Vey")
    st, out = call("PUT", "/fields", {"target": f"character:{moki}", "fields": {"height": "1.68 m"}})
    assert st == 200 and out["fields"]["height"]["value"] == "1.68 m"
    st, got = call("GET", "/fields", query={"target": [f"character:{moki}"]})
    assert got["fields"]["height"]["status"] == "supplied"
    st, out = call("POST", "/fields/approve", {"all_targets": True})
    assert st == 200 and out["ok"]
    st, j = call("POST", "/fill", {"ai": False})
    assert st == 202 and forge.wait(app, "author", j["job_id"])["status"] == "done"
    st, dx = call("GET", "/export", query={"format": ["docx"]})
    assert st == 200 and dx[0] == "__binary__" and dx[2][:2] == b"PK"
    st, hx = call("GET", "/export", query={"format": ["html"]})
    assert hx[2].startswith(b"<!doctype html>")
    with pytest.raises(AkoweError) as e:
        call("GET", "/export", query={"format": ["pdf"]})
    assert e.value.code == "CAPABILITY_UNAVAILABLE" and e.value.errors[0]["problem"] == "RENDER_NOT_WIRED"
    with pytest.raises(AkoweError):
        call("GET", "/export", query={"format": ["odt"]})
