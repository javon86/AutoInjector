"""Author driver for desktop-app/test/verify-acceptance.test.js: a person using the Akọwe web UI (Playwright,
Chromium) while Akọwe talks to the AIs through the REAL AutoInjector bridge.

ACCEPT_INFO (json from the harness): bf_url, token, control (harness control server), shots.
Phases (each timed): 0 pair through #/pair · 1 wizard + guided path + downloads + file checks · 2 ⚡ Make a book for me
· 3 retitle -> format again -> export again (+ title page check) · archive · backup · change · restore · checks ·
4 a refused outline: Redo with AI / Edit · 5 layout at 1366 / 390 px.
Every check prints ok/FAIL; the exit code is the number of failed checks (0 = PASS).
"""
from __future__ import annotations

import io
import json
import os
import re
import subprocess
import sys
import time
import urllib.request
import zipfile
from pathlib import Path

from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))
from akowe.exporters.book import parse_chapter  # noqa: E402

INFO = json.loads(os.environ["ACCEPT_INFO"])
SHOTS = Path(INFO["shots"])
SHOTS.mkdir(parents=True, exist_ok=True)
DL = SHOTS / "downloads"
CHROMIUM = os.environ.get("BF_CHROMIUM", "/opt/pw-browsers/chromium")
EPUBCHECK = os.environ.get("BF_EPUBCHECK_JAR") or str(Path(os.environ.get("SCRATCH", "/nonexistent")) / "epubcheck.jar")
OP = urllib.request.build_opener(urllib.request.ProxyHandler({}))
FAILS: list[str] = []
TIMES: list[tuple[str, float]] = []
CODE_RE = re.compile(r"\b[A-Z]{3,}(?:_[A-Z]+)+\b")
BAD_RE = re.compile(r"\bundefined\b|\[object Object\]|\bNaN\b")


def check(cid, ok, msg):
    print(f"  {'ok  ' if ok else 'FAIL'} [{cid}] {msg}", flush=True)
    if not ok:
        FAILS.append(cid)
    return ok


def api(method, path, body=None):
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(INFO["bf_url"] + "api" + path, data=data, method=method)
    req.add_header("Authorization", "Bearer " + INFO["token"])
    if data:
        req.add_header("Content-Type", "application/json")
    with OP.open(req, timeout=60) as r:
        return json.loads(r.read() or b"null")


def ctl(p):
    return json.load(OP.open(INFO["control"] + p, timeout=60))


class Phase:
    def __init__(self, name):
        self.name = name

    def __enter__(self):
        self.t = time.time()
        print(f"\n== {self.name}", flush=True)

    def __exit__(self, *a):
        TIMES.append((self.name, time.time() - self.t))
        print(f"   ({time.time() - self.t:.1f} s)", flush=True)


def shot(page, name, full=True):
    page.screenshot(path=str(SHOTS / f"{name}.png"), full_page=full)


def problems_of(page):
    out = list(page.bf_problems)
    page.bf_problems.clear()
    return out


def visible_text_ok(page, where):
    txt = page.inner_text("body")
    codes = sorted(set(CODE_RE.findall(txt)) - {"ANTHROPIC_API_KEY"})
    check(f"text:{where}", not BAD_RE.search(txt) and not codes,
          f"{where}: no undefined/NaN and no error codes shown ({codes[:5]})")


def book_top(page):
    top = page.inner_text("#main .book-top")
    m = re.search(r"(\d+) of (\d+) steps done", top)
    n = re.search(r"Do the next step: ([^\n]+)", top)
    return int(m.group(1)), int(m.group(2)), (n.group(1).strip() if n else None), "Your book is made" in top


def api_book(pid):
    b = api("GET", f"/projects/{pid}/branches/main/book")
    st = next((s for s in b["steps"] if s["id"] == b["next"]), None)
    return b["done_count"], b["total"], (st["title"] if st else None), b


def goto(page, h):
    page.evaluate("h => { location.hash = h; }", h)
    page.wait_for_timeout(1200)


def wait_msg(page, timeout=600000):
    page.wait_for_function("""() => { const m = document.querySelector('#main .book-msg');
        return m && !m.querySelector('.state.loading, .spinner'); }""", timeout=timeout)
    page.wait_for_timeout(300)


def next_steps(page, pid, limit=40, label=""):
    """Press 'Do the next step' (and approve) until the book is made; the page must agree with the API each time."""
    msgs = []
    for _ in range(limit):
        ui, ab = book_top(page), api_book(pid)
        if (ui[0], ui[1], ui[2]) != ab[:3]:
            check(f"agree{label}", False, f"page says {ui[:3]}, Akọwe says {ab[:3]}")
            break
        if ui[3] or not ui[2]:
            break
        t0 = time.time()
        page.locator("#main .book-top button", has_text="Do the next step").click()
        page.wait_for_timeout(400)
        yes = page.get_by_role("button", name="✔ Yes, I approve")
        if yes.count():
            yes.click()
        wait_msg(page)
        msg = page.inner_text("#main .book-msg")
        msgs.append((ui[2], round(time.time() - t0, 1), msg.replace("\n", " ")[:160]))
        print(f"    {ui[0]}/{ui[1]} {ui[2]}: {time.time() - t0:.1f}s · {msg.splitlines()[-1][:100] if msg else ''}", flush=True)
        if "Failed" in msg:
            return msgs, msg
    return msgs, None


def download_all(page, scope_sel, dest: Path):
    dest.mkdir(parents=True, exist_ok=True)
    got = {}
    for f in ("PDF", "EPUB", "DOCX"):
        with page.expect_download(timeout=30000) as d:
            page.locator(f"{scope_sel} button", has_text=f).first.click()
        p = dest / d.value.suggested_filename
        d.value.save_as(p)
        got[f] = p
    return got


def expected_chapters(pid):
    b = api("GET", f"/projects/{pid}/branches/main/book")
    recs = {(r["record_id"], r["version"]): r for r in api("GET", f"/projects/{pid}/branches/main/export")["records"]}
    out = []
    for c in next(s for s in b["steps"] if s["id"] == "chapters")["chapters"]:
        pr = c.get("prose") or {}
        r = recs.get((pr.get("record_id"), pr.get("version")))
        title, blocks = parse_chapter(c["chapter_id"], r["data"]["text"])
        out.append((c["chapter_id"], title, [b["text"] for b in blocks if b["t"] in ("p", "h2")]))
    return out


def norm(s):
    return " ".join(s.replace("­", "").split())


def check_files(tag, files, pid, title, author):
    import docx  # noqa: PLC0415
    import pypdfium2 as pdfium  # noqa: PLC0415
    from lxml import html as LH  # noqa: PLC0415
    chs = expected_chapters(pid)
    want_heads = [f"Chapter {i}" + (f": {t}" if t else "") for i, (_c, t, _p) in enumerate(chs, 1)]
    paras = [norm(p) for _c, _t, ps in chs for p in ps]
    texts = {}
    # EPUB
    z = zipfile.ZipFile(files["EPUB"])
    opf_name = next(n for n in z.namelist() if n.endswith(".opf"))
    opf = z.read(opf_name).decode()
    xh = sorted(n for n in z.namelist() if n.endswith(".xhtml") and "/text/" in n)
    docs = [LH.fromstring(z.read(n)) for n in xh]
    heads = [norm(h.text_content()) for d in docs for h in d.xpath("//h1")]
    texts["EPUB"] = norm(" ".join(d.text_content() for d in docs))
    check(f"{tag}:epub-meta", f"<dc:title>{title}</dc:title>" in opf.replace(' id="title"', "") or re.search(
        rf"<dc:title[^>]*>{re.escape(title)}<", opf) is not None, f"{tag} EPUB metadata title = {title!r}")
    check(f"{tag}:epub-titlepage", norm(docs[0].text_content()).startswith(title) and author in docs[0].text_content(),
          f"{tag} EPUB title page shows {title!r} by {author!r} (got {norm(docs[0].text_content())[:80]!r})")
    check(f"{tag}:epub-heads", heads[2:] == want_heads, f"{tag} EPUB chapter headings {heads[2:]}")
    if Path(EPUBCHECK).exists():
        r = subprocess.run(["java", "-jar", EPUBCHECK, str(files["EPUB"])], capture_output=True, text=True)
        check(f"{tag}:epubcheck", r.returncode == 0, f"{tag} EPUBCheck clean ({(r.stdout + r.stderr).strip().splitlines()[-1:]})")
    # DOCX
    d = docx.Document(str(files["DOCX"]))
    ps = [(p.style.name, p.text) for p in d.paragraphs]
    texts["DOCX"] = norm(" ".join(t for _s, t in ps))
    check(f"{tag}:docx-titlepage", [t for s, t in ps if s in ("Title", "Subtitle")] == [title, author],
          f"{tag} DOCX title page {[t for s, t in ps if s in ('Title', 'Subtitle')]}")
    check(f"{tag}:docx-heads", [t for s, t in ps if s.startswith("Heading 1")] == want_heads, f"{tag} DOCX headings")
    check(f"{tag}:docx-meta", d.core_properties.title == title and d.core_properties.author == author, f"{tag} DOCX properties")
    # PDF
    pdf = pdfium.PdfDocument(str(files["PDF"]))
    pages = [pdf[i].get_textpage().get_text_range() for i in range(len(pdf))]
    for i in range(len(pdf)):
        pdf[i].render(scale=2).to_pil().save(SHOTS / f"pdf_{tag}_p{i + 1}.png")
    texts["PDF"] = norm(" ".join(pages).replace("-\r\n", "").replace("\r\n", " "))
    check(f"{tag}:pdf-titlepage", norm(pages[0]) == norm(f"{title} {author}"), f"{tag} PDF page 1 = {norm(pages[0])!r}")
    for fmt, full in texts.items():
        pos, bad = -1, []
        for p in paras:
            n = full.count(p)
            if n != 1:
                bad.append(f"{p[:40]!r} x{n}")
            else:
                i = full.index(p)
                if i < pos:
                    bad.append(f"{p[:40]!r} out of order")
                pos = i
        check(f"{tag}:{fmt.lower()}-text", not bad, f"{tag} {fmt}: every chapter paragraph once, in order ({len(paras)} paragraphs) {bad[:3]}")


def watch(page):
    page.bf_problems = []
    page.on("console", lambda m: m.type == "error" and page.bf_problems.append(f"console.error: {m.text}"))
    page.on("pageerror", lambda e: page.bf_problems.append(f"pageerror: {e}"))
    page.on("response", lambda r: r.status >= 400 and "favicon" not in r.url and page.bf_problems.append(f"HTTP {r.status} {r.url}"))
    return page


def main():
    with sync_playwright() as pw:
        browser = pw.chromium.launch(executable_path=CHROMIUM, args=["--no-sandbox", "--no-proxy-server"])
        ctx = browser.new_context(viewport={"width": 1366, "height": 900}, accept_downloads=True)
        page = watch(ctx.new_page())
        dialogs = []
        page.on("dialog", lambda d: (dialogs.append(d.message), d.accept()))

        with Phase("0 · pair Akọwe with AutoInjector (#/pair)"):
            page.goto(INFO["bf_url"] + "#token=" + INFO["token"])
            page.wait_for_selector("#main h1")
            link = ctl("/connect")["link"]
            check("pair-link", bool(link), f"AutoInjector's Connect gives a #/pair link ({link})")
            page.goto(link)
            page.get_by_role("button", name=re.compile("Connect to")).click()
            page.wait_for_selector("text=Continue to my current task", timeout=30000)
            txt = page.inner_text("#main")
            check("pair-ok", "Connected to AutoInjector" in txt and txt.count("Ready") >= 3, "the page says Connected, three AIs Ready")
            shot(page, "A0_paired")
            st = ctl("/status")["status"]["connect"]["phase"]
            check("pair-ai-side", st == "connected", f"AutoInjector says {st}")

        with Phase("1 · Start a new book (wizard) + the guided path"):
            goto(page, "#/start")
            page.fill("#wz-title", "The Lantern at Gullhaven")
            page.fill("#wz-author", "Kenneth J. Christian")
            page.fill("#wz-genre", "Mystery")
            page.select_option("#wz-audience", label="Adult")
            page.select_option("#wz-length", value="short_story")
            page.fill("#wz-tone", "Spooky, warm")
            page.fill("#wz-idea", "A lighthouse keeper vanishes; his apprentice Mara searches the fog-bound harbor town for him.")
            page.get_by_role("button", name="Next ▶").click()
            page.get_by_text("Automatic", exact=True).first.click()
            page.get_by_role("button", name="Next ▶").click()
            page.wait_for_timeout(500)
            page.get_by_role("button", name="Next ▶").click()
            page.wait_for_timeout(500)
            shot(page, "A1_wizard_review")
            page.get_by_role("button", name="🚀 Create my book").click()
            page.wait_for_url(re.compile(r"#/p/[^/]+/main/book$"), timeout=30000)
            pid1 = page.url.split("#/p/")[1].split("/")[0]
            t0 = time.time()
            while time.time() - t0 < 300:
                s = api("GET", f"/bible-builder?project_id={pid1}&branch_id=main")["session"]
                if s and not s["busy"] and s["items"]:
                    break
                time.sleep(1)
            print(f"    the wizard's characters/places were made in {time.time() - t0:.1f}s: {[i['name'] for i in s['items']]}")
            page.wait_for_timeout(2500)
            canon = page.inner_text("#step-canon")
            check("ACCEPT-3 canon-note", "still making" not in canon,
                  "the canon step stops saying 'still making…' once the AI is done (without a reload)")
            page.locator("#main .book-top button", has_text="Do the next step").click()
            dlg = page.locator("#bible-dialog")
            dlg.get_by_role("button", name="Create selected as drafts").wait_for(timeout=15000)
            page.wait_for_function("() => !document.querySelector('#bible-dialog button.primary[disabled]')", timeout=15000)
            dlg.get_by_role("button", name="Create selected as drafts").click()
            dlg.locator(".bb-created").wait_for(timeout=15000)
            page.keyboard.press("Escape")
            page.wait_for_timeout(1500)
            canon = page.inner_text("#step-canon")
            check("ACCEPT-3 after-save", "not saved yet" not in canon and "Approve" in canon,
                  "after saving in the Bible Builder the book page lists the drafts to approve (without a reload)")
            if "Approve" not in canon:
                page.reload()
                page.wait_for_selector("#step-canon")
            msgs, failed = next_steps(page, pid1)
            check("guided-done", book_top(page)[3] and not failed, f"the guided path makes the book ({len(msgs)} steps)")
            shot(page, "A1_book_made")
            files1 = download_all(page, "#step-export .downloads", DL / "book1")
            check("dl-names", {p.name for p in files1.values()} == {f"The Lantern at Gullhaven.{x}" for x in ("pdf", "epub", "docx")},
                  f"downloads named after the book ({sorted(p.name for p in files1.values())})")
            check_files("book1", files1, pid1, "The Lantern at Gullhaven", "Kenneth J. Christian")
            visible_text_ok(page, "book page")
            check("clean-1", not problems_of(page), "no console errors / failed requests")

        with Phase("2 · ⚡ Make a book for me"):
            page.locator("#main button", has_text="Make a book for me").first.click()
            d = page.locator("#mab-dialog")
            d.get_by_role("button", name="⚡ Start — let the AI choose everything").click()
            t0, last, k = time.time(), "", 0
            while time.time() - t0 < 900:
                page.wait_for_timeout(1500)
                head = d.inner_text()[:300]
                m = re.search(r"(\d+ of \d+ steps)", head)
                if m and m.group(1) != last:
                    last = m.group(1)
                    print(f"    +{time.time() - t0:.0f}s {last}", flush=True)
                    shot(page, f"A2_run_{k}", full=False)
                    k += 1
                if re.search(r"Your book is made", head) or re.search(r"Cancelled|Paused|stopped", head):
                    break
            shot(page, "A2_run_end", full=False)
            check("autobook-done", "Your book is made" in d.inner_text(), f"the run finishes ({last})")
            pid2 = next(p["project_id"] for p in api("GET", "/projects") if p["project_id"] != pid1)
            title2 = api("GET", f"/projects/{pid2}/profile")["title"]
            author2 = api("GET", f"/projects/{pid2}/profile")["author"]
            files2 = download_all(page, "#mab-dialog", DL / "book2")
            check_files("book2", files2, pid2, title2, author2)
            d.locator("button", has_text="✕ Close").click()
            check("clean-2", not problems_of(page), "no console errors / failed requests")

        with Phase("3 · retitle, format again, export again; archive; backup; restore"):
            goto(page, f"#/p/{pid1}/main/book")
            page.get_by_role("button", name="📝 Book profile").click()
            dlg = page.locator("dialog[open]")
            dlg.locator("input[id$='-title']").fill("The Last Light of Gullhaven")
            dlg.get_by_role("button", name="Save profile").click()
            page.wait_for_timeout(1500)
            fmt = page.inner_text("#step-format")
            exp = page.inner_text("#step-export")
            check("retitle-marked", "Make the format again" in fmt and "(out of date)" in exp,
                  "the format and the files are marked out of date")
            shot(page, "A3_out_of_date")
            next_steps(page, pid1, label="-3")
            files1b = download_all(page, "#step-export .downloads", DL / "book1-retitled")
            check_files("book1-retitled", files1b, pid1, "The Last Light of Gullhaven", "Kenneth J. Christian")
            goto(page, "#/")
            listed = page.inner_text("#main")
            check("ACCEPT-4 list-title", "The Last Light of Gullhaven" in listed, "the book list shows the new title")
            item = page.locator("#main li, #main .record-item").filter(has_text=title2).first
            item.get_by_role("button", name="🗄 Archive").click()
            page.wait_for_timeout(1500)
            check("archive", title2 not in page.inner_text("#main").split("🗄 Archived books")[0], f"{title2!r} archived")
            goto(page, "#/settings")
            card = page.locator("section.card", has=page.locator("#bk-h"))
            card.get_by_role("button", name="💾 Back up now").click()
            page.locator("#toast:not([hidden])").wait_for(timeout=30000)
            page.wait_for_timeout(1500)
            goto(page, f"#/p/{pid1}/main/book")
            page.get_by_role("button", name="📝 Book profile").click()
            dlg = page.locator("dialog[open]")
            dlg.locator("input[id$='-author']").fill("Changed After Backup")
            dlg.get_by_role("button", name="Save profile").click()
            page.wait_for_timeout(1000)
            goto(page, "#/settings")
            card = page.locator("section.card", has=page.locator("#bk-h"))
            card.locator("li", has_text="You pressed Back up now").first.get_by_role("button", name="↩ Restore").click()
            dlg = page.locator("dialog[open]")
            dlg.locator("input[type=checkbox]").check()
            dlg.locator("input[type=text]").fill("RESTORE")
            dlg.get_by_role("button", name="↩ Restore this backup").click()
            page.wait_for_function("() => !document.querySelector('dialog[open]')", timeout=120000)
            toast = page.inner_text("#toast")
            check("ACCEPT-6 restore-toast", not re.search(r"\d{4}-\d\d-\d\dT|bk-\d+", toast), f"restore message is plain ({toast[:120]!r})")
            p1 = api("GET", f"/projects/{pid1}/profile")
            b1, b2 = api_book(pid1)[3], api_book(pid2)[3]
            check("restore-book1", p1["title"] == "The Last Light of Gullhaven" and p1["author"] == "Kenneth J. Christian"
                  and b1["done_count"] == b1["total"], f"book 1 is back as backed up ({p1['title']} / {p1['author']}, {b1['done_count']}/{b1['total']})")
            check("restore-book2", [p["name"] for p in api("GET", "/projects?archived=1")] == [title2] and b2["done_count"] == b2["total"],
                  "book 2 is still archived and made")
            goto(page, f"#/p/{pid1}/main/book")
            files1c = download_all(page, "#step-export .downloads", DL / "book1-restored")
            same = all(files1c[f].read_bytes() == files1b[f].read_bytes() for f in files1b)
            check("restore-files", same, "the files after the restore are the files from before")
            check("clean-3", not problems_of(page), "no console errors / failed requests")

        with Phase("4 · a refused outline: Redo with AI / Edit"):
            goto(page, "#/start")
            page.fill("#wz-title", "No Harbor Here")
            page.fill("#wz-author", "K. J. C.")
            page.select_option("#wz-length", value="short_story")
            page.fill("#wz-idea", "A cartographer maps a city that rearranges itself every night.")
            page.get_by_role("button", name="Next ▶").click()
            page.get_by_text("Automatic", exact=True).first.click()
            page.get_by_role("button", name="Next ▶").click()
            nums = page.locator("#main input.wz-num")
            for i, v in enumerate((1, 0, 0)):
                nums.nth(i).fill(str(v))
                nums.nth(i).dispatch_event("change")
            page.get_by_role("button", name="Next ▶").click()
            page.get_by_role("button", name="🚀 Create my book").click()
            page.wait_for_url(re.compile(r"#/p/[^/]+/main/book$"), timeout=30000)
            pid3 = page.url.split("#/p/")[1].split("/")[0]
            t0 = time.time()
            while time.time() - t0 < 300:
                s = api("GET", f"/bible-builder?project_id={pid3}&branch_id=main")["session"]
                if s and not s["busy"] and s["items"]:
                    break
                time.sleep(1)
            page.reload()
            page.wait_for_selector("#step-canon")
            page.locator("#main .book-top button", has_text="Do the next step").click()
            dlg = page.locator("#bible-dialog")
            dlg.get_by_role("button", name="Create selected as drafts").wait_for(timeout=15000)
            page.wait_for_function("() => !document.querySelector('#bible-dialog button.primary[disabled]')", timeout=15000)
            dlg.get_by_role("button", name="Create selected as drafts").click()
            dlg.locator(".bb-created").wait_for(timeout=15000)
            page.keyboard.press("Escape")
            page.reload()
            page.wait_for_selector("#step-canon")
            _msgs, failed = next_steps(page, pid3, label="-4")
            shot(page, "A4_refused")
            check("refused-shown", bool(failed) and "Redo with AI" in failed and "Edit it" in failed,
                  "a refused outline offers Redo with AI and Edit")
            if failed:
                check("ACCEPT-7 refused-plain", not CODE_RE.search(failed), f"the refusal is in plain words ({failed.splitlines()[2][:120]!r})")
                page.locator("#main .book-msg button", has_text="Redo with AI").first.click()
                page.wait_for_timeout(500)
                wait_msg(page)
                shot(page, "A4_after_redo")
                print("    after Redo:", page.inner_text("#step-outline").replace("\n", " ")[:200])

        with Phase("5 · layout at 1366 and 390 px"):
            for w, hgt in ((1366, 900), (390, 844)):
                page.set_viewport_size({"width": w, "height": hgt})
                for r in ("#/", "#/start", f"#/p/{pid1}/main/book", f"#/p/{pid1}/main/review", "#/shelf", "#/settings"):
                    goto(page, r)
                    ov = page.evaluate("document.documentElement.scrollWidth - document.documentElement.clientWidth")
                    check(f"overflow-{w}{r}", ov <= 0, f"{w}px {r}: no sideways scroll ({ov})")
                    shot(page, f"A5_{w}_{re.sub(r'[^a-z0-9]+', '_', r)}", full=(w == 390))
            visible_text_ok(page, "settings")
        browser.close()

    rep = ctl("/status")
    print(f"\nAutoInjector: phase={rep['status']['connect']['phase']} · posts {json.dumps(rep['dumps'])} · Akọwe prompts {rep['hook']['bf']}")
    print("Timings: " + " · ".join(f"{n.split(' · ')[0]} {s:.0f}s" for n, s in TIMES))
    print(f"\n{len(FAILS)} check(s) failed: {', '.join(FAILS)}" if FAILS else "\nall checks passed")
    return len(FAILS)


if __name__ == "__main__":
    sys.exit(min(main(), 100))
